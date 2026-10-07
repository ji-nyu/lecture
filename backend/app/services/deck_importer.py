"""Turn an instructor-made PPTX into a lecture plan, slide spec, and presentation file.

The uploaded deck is the presentation. Scripts and video are written against those slides.
"""

from __future__ import annotations

import hashlib
import json
import logging
import shutil
import uuid
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from ..errors import AnalysisNotReady, DeckNeedsPptx, EmptyDeck, ProfileNotReady, SourceNotUploaded
from ..models.lecture_plan import (
    LecturePlan,
    LectureSection,
    PlanMetrics,
    PresentationHints,
    SectionKind,
)
from ..models.lecture_profile import LectureProfile
from ..models.presentation import JobState, PresentationJob, PresentationResult
from ..models.project import PresentationStatus, Project
from ..models.slide_spec import (
    ContentOrigin,
    Slide,
    SlideMetrics,
    SlideSpecification,
    SlideStyle,
    SlideType,
)
from ..models.source import SourceAnalysis, SourceLocation, SourceMaterial
from ..storage.project_store import ProjectStore, utcnow
from .lecture_script_builder import _BAD_CONCEPT, _is_chrome_line, speakable_chars
from .lecture_script_writer import clear_script_cache
from .prompt_builder import _is_structural_message

logger = logging.getLogger("ailecturegen")

PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
PLANNER = "imported-deck-v1"
SECTION_ID = "sec01"


@dataclass
class DeckPage:
    number: int
    title: str
    points: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


def apply_imported_deck(store: ProjectStore, project: Project) -> Project:
    """Build plan + slides from the uploaded PPTX and register that file as the presentation."""
    if not project.source_file_path or not project.source_file:
        raise SourceNotUploaded()
    src = store.project_dir(project.id) / project.source_file_path
    if src.suffix.lower() != ".pptx" or not src.is_file():
        raise DeckNeedsPptx()
    if project.source_analysis is None:
        raise AnalysisNotReady()
    if project.lecture_profile is None:
        raise ProfileNotReady()
    material = _load_material(store, project.id)
    pages = extract_pages(src)
    plan, spec = build_plan_and_spec(project, pages, material)
    project.lecture_plan = plan
    project.slide_specification = spec
    project.enriched_specification = None
    clear_script_cache(store, project.id)
    project.discard_video()
    _install_presentation(store, project, src, spec.slide_count)
    project.presentation_status = PresentationStatus.completed
    project.error_message = None
    store.save(project)
    return project


def extract_pages(path: Path) -> list[DeckPage]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    pages: list[DeckPage] = []
    for idx, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = ""
        title_id = None
        if title_shape is not None:
            title = (title_shape.text_frame.text or "").strip()
            title_id = title_shape.shape_id
        points: list[str] = []
        for shp in _walk(slide.shapes):
            if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
                continue
            if not getattr(shp, "has_text_frame", False) or not shp.has_text_frame:
                continue
            if title_id is not None and shp.shape_id == title_id:
                continue
            for para in shp.text_frame.paragraphs:
                text = "".join(r.text for r in para.runs).strip()
                if len(text) >= 2:
                    points.append(text)
        if not title:
            title = next((p[:60] for p in points if p), f"슬라이드 {idx}")
        notes: list[str] = []
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            for ln in (slide.notes_slide.notes_text_frame.text or "").splitlines():
                t = ln.strip()
                if len(t) >= 8 and not _is_chrome_line(t) and not _is_structural_message(t):
                    notes.append(t)
        pages.append(DeckPage(number=idx, title=title or f"슬라이드 {idx}", points=points, notes=notes))
    if not pages:
        raise EmptyDeck()
    return pages


def build_plan_and_spec(
    project: Project,
    pages: list[DeckPage],
    material: SourceMaterial | None,
) -> tuple[LecturePlan, SlideSpecification]:
    profile = project.lecture_profile
    analysis = project.source_analysis
    assert profile is not None and analysis is not None
    if not pages:
        raise EmptyDeck()
    n = len(pages)
    minutes = profile.duration_minutes
    times = _allocate_speakable([_weight(p) for p in pages], minutes * 60)
    by_page = {s.page: s for s in (material.sections if material else []) if s.page}
    slides: list[Slide] = []
    source_points: list[str] = []
    for page, seconds in zip(pages, times):
        sec = by_page.get(page.number)
        points = [p for p in page.points if p and p != page.title and not _is_chrome_line(p)]
        if page.notes:
            points.extend(n for n in page.notes if not _is_chrome_line(n) and not _is_structural_message(n))
        if not points and page.title and not _is_chrome_line(page.title):
            points = [page.title]
        message = next((p for p in points if len(p) >= 8 and not _is_structural_message(p)), page.title)
        stype = _guess_type(page, page.number, n)
        sid = sec.id if sec else f"s{page.number:02d}"
        slides.append(
            Slide(
                slide_number=page.number,
                section_id=SECTION_ID,
                section_title="업로드한 슬라이드",
                title=page.title,
                slide_type=stype,
                learning_purpose=f"업로드한 슬라이드 ‘{page.title}’의 내용을 강의로 설명한다.",
                key_message=message,
                message_from_source=bool(points),
                key_points=points[:40],
                source_reference=SourceLocation(
                    section_id=sid,
                    section_title=sec.title if sec else page.title,
                    line=None,
                    page=page.number,
                ),
                visual_instruction="업로드한 슬라이드 화면을 그대로 사용한다.",
                presenter_instruction="화면에 있는 문장을 강의 말로 설명한다.",
                estimated_explanation_time=seconds,
                concepts=_concepts_on_slide(analysis, page),
                content_origin=ContentOrigin.source,
            )
        )
        source_points.extend(points[:3])
    for slide, seconds in zip(slides, _allocate_speakable([speakable_chars(s) for s in slides], minutes * 60)):
        slide.estimated_explanation_time = seconds
    titles = [p.title for p in pages if p.title and not p.title.startswith("슬라이드 ")]
    objectives = [f"‘{t}’을 설명할 수 있다." for t in titles[:3]]
    if not objectives:
        objectives = ["업로드한 슬라이드의 내용을 설명할 수 있다."]
    now = datetime.now(timezone.utc)
    section = LectureSection(
        id=SECTION_ID,
        order=1,
        title="업로드한 슬라이드",
        purpose="교수자가 만든 PPT 순서대로 설명한다.",
        kind=SectionKind.concept,
        duration_minutes=minutes,
        importance=5,
        concepts=list(dict.fromkeys(c for s in slides for c in s.concepts))[:12],
        source_points=source_points[:12],
        source_section_ids=[],
        teaching_notes=["업로드한 PPT 장 구성을 바꾸지 않습니다."],
        estimated_slides=n,
    )
    plan = LecturePlan(
        id=uuid.uuid4().hex,
        project_id=project.id,
        title=project.title or analysis.title or "업로드한 강의",
        audience=profile.audience_level.value,
        duration_minutes=minutes,
        difficulty=profile.difficulty.value,
        lecture_type=profile.lecture_type.value,
        explanation_depth=profile.explanation_depth.value,
        source_policy=profile.source_policy.value,
        learning_objectives=objectives,
        estimated_slide_count=n,
        sections=[section],
        metrics=PlanMetrics(
            section_count=1,
            slide_count=n,
            concept_count=len(section.concepts),
            definition_count=sum(1 for s in slides if s.slide_type == SlideType.definition),
            example_count=sum(1 for s in slides if s.slide_type == SlideType.example),
            practice_activity_count=0,
            practice_step_count=0,
            code_count=0,
            quiz_question_count=0,
            explanation_minutes=minutes,
            practice_minutes=0,
            minutes_by_kind={"concept": minutes},
        ),
        presentation_hints=PresentationHints(
            lecture_tone=profile.lecture_tone.value,
            speaker_notes=profile.speaker_notes.value,
            visual_level=profile.visual_level.value,
            slide_density=profile.slide_density.value,
            source_policy=profile.source_policy.value,
        ),
        warnings=["업로드한 PPT를 그대로 사용합니다. 슬라이드를 새로 만들지 않습니다."],
        profile_id=profile.id,
        source_id=analysis.source_id,
        planner=PLANNER,
        generated_at=now,
    )
    by_type: dict[str, int] = {}
    for s in slides:
        by_type[s.slide_type.value] = by_type.get(s.slide_type.value, 0) + 1
    pts = [len(s.key_points) for s in slides]
    spec = SlideSpecification(
        lecture_id=plan.id,
        project_id=project.id,
        title=plan.title,
        duration_minutes=minutes,
        slide_count=n,
        plan_estimated_slide_count=n,
        slides=slides,
        metrics=SlideMetrics(
            slide_count=n,
            total_seconds=minutes * 60,
            slides_by_type=by_type,
            slides_by_section={SECTION_ID: n},
            source_slide_count=n,
            suggested_slide_count=0,
            avg_key_points=round(sum(pts) / len(pts), 2) if pts else 0,
            max_key_points=max(pts) if pts else 0,
            visual_slide_count=0,
        ),
        style=SlideStyle(
            slide_density=profile.slide_density.value,
            visual_level=profile.visual_level.value,
            lecture_tone=profile.lecture_tone.value,
            speaker_notes=profile.speaker_notes.value,
            source_policy=profile.source_policy.value,
            audience=profile.audience_level.value,
            difficulty=profile.difficulty.value,
            max_key_points=8,
        ),
        warnings=list(plan.warnings),
        planner=PLANNER,
        generated_at=now,
    )
    return plan, spec


def _install_presentation(store: ProjectStore, project: Project, src: Path, slide_count: int) -> None:
    dest_dir = store.presentation_dir(project.id)
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / "presentation.pptx"
    if src.resolve() != dest.resolve():
        shutil.copyfile(src, dest)
    data = dest.read_bytes()
    now = utcnow()
    job_id = f"imp-{uuid.uuid4().hex}"
    name = project.source_file.filename if project.source_file else "lecture.pptx"
    if not name.lower().endswith(".pptx"):
        name = Path(name).stem + ".pptx"
    project.presentation_provider = "imported"
    project.presentation_job = PresentationJob(
        provider="imported",
        is_mock=False,
        job_id=job_id,
        state=JobState.completed,
        progress=100,
        stage="업로드한 PPT를 사용합니다.",
        prompt_hash="imported-deck",
        created_at=now,
        updated_at=now,
    )
    project.presentation_result = PresentationResult(
        provider="imported",
        is_mock=False,
        job_id=job_id,
        prompt_hash="imported-deck",
        file_name=name,
        content_type=PPTX_TYPE,
        size_bytes=len(data),
        sha256=hashlib.sha256(data).hexdigest(),
        slide_count=slide_count,
        completed_at=now,
        note="교수자가 업로드한 PPT를 그대로 사용합니다.",
    )


def _load_material(store: ProjectStore, project_id: str) -> SourceMaterial | None:
    path = store.project_dir(project_id) / "source_material.json"
    try:
        if path.is_file():
            return SourceMaterial.model_validate(json.loads(path.read_text(encoding="utf-8")))
    except (OSError, ValueError):
        logger.warning("source_material.json could not be read for imported deck")
    return None


def _walk(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shp in shapes:
        if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shp.shapes)
        else:
            yield shp


def _weight(page: DeckPage) -> int:
    from .lecture_script_builder import _usable_line

    n = 0
    for raw in (page.title, *page.points, *page.notes):
        line = _usable_line(raw)
        if line:
            n += len(line)
    return max(20, n)


def _allocate_speakable(weights: list[int], total: int) -> list[int]:
    """Give each slide up to ~5× its speakable length so the script can fill the slot."""
    from .lecture_script_builder import SCRIPT_CHARS_PER_SECOND

    times = _allocate(weights, total)
    caps = [max(8, int(max(1, w) / SCRIPT_CHARS_PER_SECOND * 5) + 8) for w in weights]
    leftover = 0
    out: list[int] = []
    for t, cap in zip(times, caps):
        if t > cap:
            leftover += t - cap
            out.append(cap)
        else:
            out.append(t)
    n = len(out)
    i = 0
    guard = 0
    while leftover > 0 and n and guard < n * (leftover + 3):
        room = caps[i % n] - out[i % n]
        if room > 0:
            take = min(room, leftover)
            out[i % n] += take
            leftover -= take
        i += 1
        guard += 1
    if leftover and out:
        order = sorted(range(n), key=lambda i: weights[i], reverse=True)
        out[order[0]] += leftover
    return out or times


def _allocate(weights: list[int], total: int) -> list[int]:
    n = len(weights)
    floor = 8
    if n * floor > total:
        floor = max(4, total // n)
    leftover = max(0, total - floor * n)
    wsum = sum(max(1, w) for w in weights) or n
    extras = [leftover * max(1, w) // wsum for w in weights]
    times = [floor + e for e in extras]
    drift = total - sum(times)
    i = 0
    while drift != 0 and times:
        step = 1 if drift > 0 else -1
        if times[i % n] + step >= 1:
            times[i % n] += step
            drift -= step
        i += 1
        if i > n * (abs(drift) + 2):
            break
    if times and sum(times) != total:
        times[-1] += total - sum(times)
        if times[-1] < 1:
            times[-1] = 1
    return times


def _guess_type(page: DeckPage, number: int, total: int) -> SlideType:
    t = page.title
    blob = " ".join([page.title, *page.points[:4]])
    if number == 1 and total > 1 and len(page.points) < 2:
        return SlideType.title
    if any(k in t for k in ("목차", "학습 목표", "강의 구성", "차례")):
        return SlideType.agenda
    if any(k in blob for k in ("정의", "용어", "DEFINITION")):
        return SlideType.definition
    if any(k in blob for k in ("예시", "사례", "FOR EXAMPLE")):
        return SlideType.example
    if any(k in t for k in ("정리", "요약", "핵심")):
        return SlideType.summary
    if number == total and total > 2:
        return SlideType.summary
    return SlideType.concept


def _concepts_on_slide(analysis: SourceAnalysis, page: DeckPage) -> list[str]:
    blob = " ".join([page.title, *page.points])
    out: list[str] = []
    for c in analysis.concepts:
        if not c.name or c.name.upper() in _BAD_CONCEPT:
            continue
        if c.name in blob and c.name not in out:
            out.append(c.name)
        if len(out) >= 4:
            break
    return out
