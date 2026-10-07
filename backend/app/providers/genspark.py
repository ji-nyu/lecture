"""GensparkProvider (STAGE 8): renders the presentation with the official Genspark CLI (`gsk`).

All Genspark-specific code lives in this file and nowhere else; the Lecture Engine only knows
`PresentationProvider`.

What was verified before writing this (MASTER_SPEC STAGE 8: no guessed endpoints, parameters or
response schemas):
  * Package `@genspark/cli` (npm; installed with `npm install -g @genspark/cli`, binary `gsk`).
    Checked against the published package, versions 1.9.1 and 1.14.0: its README, the shipped skill
    documents (`gsk-shared`, `gsk-create-task`, `gsk-task-status`, `gsk-task-artifacts`), its source,
    and by running the real binary (`--version`, `--help`, `task --help`, `task export --help`).
  * Authentication: `GSK_API_KEY` environment variable (or `--api-key`, or `gsk login`, which saves the
    credentials in `~/.genspark-tool-cli/config.json` or `GSK_CONFIG`); API base URL
    `--base-url` / `GSK_BASE_URL` (default https://www.genspark.ai); `GSK_NO_AUTO_UPDATE=1` disables
    the CLI's background self-update; `--no-input` never prompts.
  * `gsk task create slides` needs `--task_name`, `--query`, `--instructions`. `--args-file <json>`
    carries those parameters without shell quoting (file keys are the parameter names).
  * VERIFIED WITH A REAL RUN (logged in with `gsk login`): `task create` only SUBMITS and returns in
    seconds (`data.run_id`, `data.project_id`, `data.task_url`); the work then runs on Genspark's server
    for minutes. `gsk task status <run_id>` answers `{"status":"ok","data":{"state":"running",
    "project_id":...}}`. `gsk task export <project_id> --format pptx -o <file>` downloads the deck.
    `--follow` (SSE stream) is NOT used: through Cloudflare it returned an HTML challenge page instead
    of events and the CLI exited 1 although the task kept running (and was billed).
  * Not seen yet: the exact name of the terminal state. The runner therefore polls while the state is
    a known "active" one, fails on a known failure state, and for any other state tries the export;
    a job counts as done ONLY when a readable .pptx exists (the exit code is not trusted either: the
    Windows build of the CLI was seen crashing at exit, code 3221226505).

How a job runs (the work takes minutes, the API must not block):
  create_presentation writes the parameters to `<jobs>/<job_id>/args.json` and starts a detached runner
  (`python -m app.providers.genspark <job_dir>`) that submits the task, polls its status, exports the
  deck, writes logs, and finally `result.json`. get_status only reads those files, so a job survives a
  server restart. A heartbeat file tells a dead runner from a slow task. No progress is available
  (progress stays 0 while running).

Security: the API key is handed to the CLI through its environment only (never a command line, never a
file), logs are redacted, and no message returned to a client contains CLI output.
"""

from __future__ import annotations

import json
import logging
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import uuid
from pathlib import Path

from ..errors import PresentationGenerationFailed, PresentationProviderUnavailable
from ..models.presentation import (
    DownloadedFile,
    JobState,
    PresentationRequest,
    ProviderResult,
    ProviderStatus,
)
from .base import PresentationProvider, ProviderAvailability

logger = logging.getLogger("ailecturegen")

PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
OUTPUT_NAME = "output.pptx"
FILE_NAME = "genspark_presentation.pptx"
DEFAULT_CLI = "gsk"
INSTALL_HINT = "npm install -g @genspark/cli"

_BACKEND_DIR = Path(__file__).resolve().parents[2]  # the directory that contains the `app` package
_JOB_RE = re.compile(r"^gsk-[0-9a-f]{32}$")  # also blocks path traversal
_ID_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._:\-]{0,199}$")  # a value we pass back to the CLI
_URL_RE = re.compile(r"^https?://[^\s]+$")
POLL_SECONDS = 20  # between two `gsk task status` calls (Genspark itself suggests 15-60 s)
_POLL_MAX_FAILURES = 8  # consecutive failed status calls (network, Cloudflare...) before giving up
_CALL_TIMEOUT_SECONDS = 180  # one status/export call must not hang
_ACTIVE_STATES = {"queued", "pending", "submitted", "starting", "running", "in_progress", "processing"}
_FAILED_STATES = {"failed", "error", "errored", "cancelled", "canceled", "stopped", "aborted", "timeout", "timed_out"}
_HEARTBEAT_SECONDS = 5
_HEARTBEAT_STALE_SECONDS = 60  # no heartbeat and no result for this long: the runner is gone
_SECRET_ENV = ("LLM_API_KEY", "OPENAI_API_KEY")  # not for the CLI

MSG_NO_CLI = f"Genspark CLI(gsk)를 찾을 수 없습니다. `{INSTALL_HINT}` 로 설치하거나 GENSPARK_CLI_PATH를 설정해 주세요."
MSG_NO_KEY = (
    "Genspark 인증 정보가 없습니다. 터미널에서 `gsk login`을 실행하거나, "
    "backend/.env에 GENSPARK_API_KEY를 설정해 주세요."
)
MSG_BAD_URL = "GENSPARK_API_URL은 http:// 또는 https:// 로 시작하는 주소여야 합니다."


def build_instructions(request: PresentationRequest) -> str:
    """The CLI's `--instructions` (the task agent's working rules): only what the spec gives
    Genspark to do and not to do, plus the fixed slide list. The lecture itself is the `query`."""
    titles = "\n".join(f"{s.slide_number}. {s.title}" for s in request.slides)
    return (
        "이 요청의 query는 이미 확정된 강의 명세입니다. 당신은 그것을 프레젠테이션으로 렌더링하는 역할만 맡습니다.\n"
        f"- 슬라이드는 정확히 {request.slide_count}장이며, 아래 목록의 순서와 제목을 그대로 따릅니다. "
        "슬라이드를 추가, 삭제, 병합, 분할하지 않습니다.\n"
        "- 맡는 일: 시각 레이아웃, 슬라이드 렌더링, 이미지 선택, 프레젠테이션 스타일링.\n"
        "- 맡지 않는 일: 강의 구조, 중요 개념, 섹션 시간, 난이도를 정하거나 바꾸는 것. "
        "명세에 적힌 텍스트와 수치를 그대로 사용하고 새로운 사실이나 예시를 덧붙이지 않습니다.\n"
        "- 각 슬라이드의 본문 항목은 화면에 읽히는 글자로 넣습니다. "
        "제목이나 개념 이름만 적힌 빈 상자·빈 카드를 만들지 않으며, 장식 이미지가 본문을 대체하지 않습니다. "
        "'설명한다'처럼 강의 진행을 말하는 문장은 화면 본문이 아닙니다.\n"
        "- 결과는 PPTX로 내보낼 수 있는 편집 가능한 슬라이드 덱이어야 합니다.\n\n"
        f"슬라이드 목록:\n{titles}"
    )


def build_query(request: PresentationRequest) -> str:
    """The CLI's `--query`: the whole prompt verbatim (the CLI docs ask for the full material, not a
    summary) with a closing sentence that states the deliverable."""
    return (
        f"{request.final_prompt}\n\n---\n"
        f"위 명세에 따라 정확히 {request.slide_count}장의 프레젠테이션(슬라이드 덱)을 만들어 주세요."
    )


class GensparkProvider(PresentationProvider):
    name = "genspark"
    is_mock = False

    def __init__(
        self,
        jobs_root: Path,
        *,
        api_key: str | None = None,
        api_url: str | None = None,
        cli_command: str | None = None,
        timeout_seconds: int = 2100,
        login_config_path: Path | None = None,
        poll_seconds: float | None = None,
    ):
        self.root = Path(jobs_root)
        self._api_key = api_key or None  # never logged, never written to disk
        # Where `gsk login` keeps its credentials (documented: GSK_CONFIG, default
        # ~/.genspark-tool-cli/config.json). Only its EXISTENCE is checked; the file is never read.
        self._login_config = login_config_path
        self.api_url = api_url or None
        self.cli_command = cli_command or DEFAULT_CLI
        self.timeout_seconds = int(timeout_seconds)
        self.poll_seconds = float(poll_seconds if poll_seconds is not None else POLL_SECONDS)

    def __repr__(self) -> str:  # the key must never show up in a log line
        return f"GensparkProvider(cli={self.cli_command!r}, key={'set' if self._api_key else 'unset'})"

    # ---- interface --------------------------------------------------------------------
    def availability(self) -> ProviderAvailability:
        if self._cli() is None:
            return ProviderAvailability(configured=False, message=MSG_NO_CLI)
        if not self._api_key and not self._has_login():
            return ProviderAvailability(configured=False, message=MSG_NO_KEY)
        if self.api_url and not _URL_RE.match(self.api_url):
            return ProviderAvailability(configured=False, message=MSG_BAD_URL)
        return ProviderAvailability(configured=True)

    def create_presentation(self, request: PresentationRequest) -> ProviderStatus:
        availability = self.availability()
        if not availability.configured:
            raise PresentationProviderUnavailable(availability.message)
        cli = self._cli()
        job_id = f"gsk-{uuid.uuid4().hex}"
        job_dir = self.root / job_id
        base = [cli, "--no-input"] + (["--base-url", self.api_url] if self.api_url else [])
        try:
            job_dir.mkdir(parents=True, exist_ok=True)
            _write_text(job_dir / "args.json", json.dumps(
                {
                    "task_name": (request.title or "강의")[:100],
                    "query": build_query(request),
                    "instructions": build_instructions(request),
                },
                ensure_ascii=False,
            ))
            _write_text(job_dir / "job.json", json.dumps(
                {
                    "job_id": job_id,
                    "created_at": time.time(),
                    "slide_count": request.slide_count,
                    "prompt_hash": request.prompt_hash,
                    "timeout_seconds": self.timeout_seconds,
                    # No key in here: the runner gets it through its environment.
                    # Submit only (no --follow / no -o): --follow goes through Cloudflare SSE and
                    # was seen returning an HTML challenge while the billed task kept running.
                    "create_command": base + [
                        "task", "create", "slides", "--args-file", "args.json",
                    ],
                    "status_command": base + ["task", "status"],
                    "export_command": base + ["task", "export"],
                    "poll_seconds": self.poll_seconds,
                }
            ))
            self._launch_runner(job_dir)
        except PresentationProviderUnavailable:
            raise
        except Exception as exc:
            logger.exception("Could not start the Genspark runner", exc_info=exc)  # server log only
            raise PresentationProviderUnavailable("Genspark 작업을 시작하지 못했습니다. 서버 로그를 확인해 주세요.") from exc
        return ProviderStatus(job_id=job_id, state=JobState.queued, stage="Genspark 작업 시작")

    def get_status(self, job_id: str) -> ProviderStatus:
        job_dir, job = self._job(job_id)
        state, message, _ = self._evaluate(job_dir, job)
        if state == JobState.completed:
            return ProviderStatus(job_id=job_id, state=state, progress=100, stage="완료")
        if state == JobState.failed:
            return ProviderStatus(job_id=job_id, state=state, stage="실패", message=message)
        started = (job_dir / "heartbeat.txt").is_file()
        return ProviderStatus(
            job_id=job_id,
            state=JobState.running if started else JobState.queued,
            stage="Genspark에서 생성 중 (진행률은 제공되지 않음)" if started else "Genspark 작업 시작",
        )

    def get_result(self, job_id: str) -> ProviderResult:
        job_dir, job = self._job(job_id)
        state, message, slides = self._evaluate(job_dir, job)
        if state != JobState.completed:
            raise PresentationGenerationFailed(message or "프레젠테이션이 아직 완성되지 않았습니다.")
        return ProviderResult(
            job_id=job_id, file_name=FILE_NAME, content_type=PPTX_TYPE,
            size_bytes=(job_dir / OUTPUT_NAME).stat().st_size, slide_count=slides,
        )

    def download(self, job_id: str) -> DownloadedFile:
        job_dir, job = self._job(job_id)
        state, message, _ = self._evaluate(job_dir, job)
        if state != JobState.completed:
            raise PresentationGenerationFailed(message or "프레젠테이션이 아직 완성되지 않았습니다.")
        return DownloadedFile(
            file_name=FILE_NAME, content_type=PPTX_TYPE, data=(job_dir / OUTPUT_NAME).read_bytes()
        )

    # ---- internals --------------------------------------------------------------------
    def _cli(self) -> str | None:
        cmd = self.cli_command
        # An explicit path (tests, GENSPARK_CLI_PATH) is never replaced by a `gsk` found elsewhere.
        if os.path.dirname(cmd) or Path(cmd).suffix:
            return str(Path(cmd)) if Path(cmd).is_file() else shutil.which(cmd)
        found = shutil.which(cmd)
        if not found and os.name == "nt":
            found = shutil.which(f"{cmd}.cmd")
        return found or _npm_global_binary(cmd)

    def _has_login(self) -> bool:
        """Did the user run `gsk login`? The CLI then finds its own credentials; we only check that
        the file exists (never open it)."""
        path = self._login_config
        if path is None:
            env = os.environ.get("GSK_CONFIG")
            path = Path(env) if env else Path.home() / ".genspark-tool-cli" / "config.json"
        try:
            return Path(path).is_file()
        except OSError:
            return False

    def _job(self, job_id: str) -> tuple[Path, dict]:
        if not _JOB_RE.match(job_id or ""):
            raise PresentationGenerationFailed("알 수 없는 프레젠테이션 작업입니다.")
        job_dir = self.root / job_id
        try:
            return job_dir, json.loads((job_dir / "job.json").read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            raise PresentationGenerationFailed("프레젠테이션 작업 정보를 찾을 수 없습니다.") from exc

    def _child_env(self) -> dict[str, str]:
        env = {k: v for k, v in os.environ.items() if k not in _SECRET_ENV}
        if self._api_key:  # otherwise the CLI uses the credentials `gsk login` saved
            env["GSK_API_KEY"] = self._api_key
        else:
            env.pop("GSK_API_KEY", None)
        env["GSK_NO_AUTO_UPDATE"] = "1"  # no background self-update while a job runs
        return env

    def _launch_runner(self, job_dir: Path) -> None:
        kwargs: dict = dict(
            cwd=str(_BACKEND_DIR), env=self._child_env(), close_fds=True,
            stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        )
        if os.name == "nt":  # outlive this process, no console window
            kwargs["creationflags"] = subprocess.DETACHED_PROCESS | subprocess.CREATE_NEW_PROCESS_GROUP
        else:
            kwargs["start_new_session"] = True
        subprocess.Popen([sys.executable, "-m", "app.providers.genspark", str(job_dir)], **kwargs)

    def _redact(self, text: str) -> str:
        return text.replace(self._api_key, "***") if self._api_key else text

    def _log_tail(self, job_dir: Path, step: str) -> str:
        parts = []
        for stream in ("stderr", "stdout"):
            try:
                data = (job_dir / f"{step}.{stream}.txt").read_text(encoding="utf-8", errors="replace")
            except OSError:
                continue
            parts.append(f"{stream}: {data[-600:].strip()}")
        return self._redact(" | ".join(parts))

    def _evaluate(self, job_dir: Path, job: dict) -> tuple[JobState, str | None, int | None]:
        """(state, user-facing message if failed, slide count of the finished file)."""
        result = _read_json(job_dir / "result.json")
        if result is None:  # still running - unless the runner died
            beat = job_dir / "heartbeat.txt"
            try:
                last = beat.stat().st_mtime if beat.is_file() else float(job["created_at"])
            except (OSError, ValueError, KeyError):
                last = 0.0
            if time.time() - last > _HEARTBEAT_STALE_SECONDS:
                return JobState.failed, "Genspark 작업을 처리하던 프로세스가 중단되었습니다. 다시 생성해 주세요.", None
            return JobState.running, None, None

        output = job_dir / OUTPUT_NAME
        hint = "API 키, 네트워크, 크레딧을 확인하고 서버 로그를 확인해 주세요."
        if result.get("timed_out"):
            minutes = max(1, int(job.get("timeout_seconds", self.timeout_seconds)) // 60)
            return JobState.failed, f"Genspark 생성이 제한 시간({minutes}분) 안에 끝나지 않아 중단했습니다.", None
        if result.get("reported_status") == "error":
            logger.warning("Genspark CLI reported an error: %s", self._log_tail(job_dir, "create"))
            return JobState.failed, f"Genspark가 작업 실패를 알려 왔습니다. {hint}", None
        if result.get("task_state") in _FAILED_STATES:
            logger.warning("Genspark task ended in %s: %s", result.get("task_state"), self._log_tail(job_dir, "create"))
            return JobState.failed, f"Genspark 작업이 실패했습니다. {hint}", None
        # The deliverable decides, not the exit code: the CLI's Windows build was seen crashing at
        # exit (exit 3221226505, a libuv assertion) after finishing its work. A readable .pptx counts.
        if not output.is_file():
            code = result.get("create_exit")
            if code not in (0, None) and not result.get("run_id") and not result.get("project_id"):
                logger.warning("Genspark CLI failed (exit %s): %s", code, self._log_tail(job_dir, "create"))
                return JobState.failed, f"Genspark 명령이 실패했습니다(종료 코드 {code}). {hint}", None
            logger.warning(
                "Genspark finished but no PPTX was exported (project_id=%s, export_exit=%s): %s",
                result.get("project_id"), result.get("export_exit"), self._log_tail(job_dir, "create"),
            )
            return JobState.failed, "Genspark 작업은 끝났지만 PPTX 파일을 받지 못했습니다. 서버 로그를 확인해 주세요.", None
        slides = _count_slides(output)
        if slides is None:
            return JobState.failed, "Genspark가 돌려준 파일이 올바른 PPTX가 아닙니다.", None
        if result.get("create_exit") != 0:
            logger.warning("Genspark CLI exited with %s but delivered a valid PPTX; using it", result.get("create_exit"))
        return JobState.completed, None, slides


# ---------------------------------------------------------------------------------------
# helpers shared with the runner
def _write_text(path: Path, text: str) -> None:
    tmp = path.with_name(path.name + ".tmp")
    tmp.write_text(text, encoding="utf-8")
    os.replace(tmp, path)


def _read_json(path: Path) -> dict | None:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    return data if isinstance(data, dict) else None


def _count_slides(path: Path) -> int | None:
    """Number of slides if `path` is a readable presentation, else None."""
    try:
        from pptx import Presentation

        return len(Presentation(str(path)).slides)
    except Exception:
        return None


def _npm_global_binary(name: str) -> str | None:
    """`gsk` after `npm install -g` when the npm global folder is not on PATH (typical for an IDE)."""
    npm = shutil.which("npm") or shutil.which("npm.cmd")
    if not npm:
        return None
    try:
        prefix = subprocess.check_output(
            [npm, "prefix", "-g"], text=True, timeout=8,
            stderr=subprocess.DEVNULL, stdin=subprocess.DEVNULL,
        ).strip()
    except (OSError, subprocess.SubprocessError):
        return None
    if not prefix:
        return None
    for candidate in (Path(prefix) / f"{name}.cmd", Path(prefix) / name, Path(prefix) / "bin" / name):
        if candidate.is_file():
            return str(candidate)
    return None


def _cli_id(value: object) -> str | None:
    return value if isinstance(value, str) and _ID_RE.match(value) else None


def _parse_cli_json(path: Path) -> dict:
    """Fields the CLI's own source reads from its JSON stdout. Anything unexpected is None."""
    empty = {"status": None, "project_id": None, "run_id": None, "state": None}
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except OSError:
        return empty
    doc = None
    for candidate in (text, text[text.find("{"): text.rfind("}") + 1] if "{" in text else ""):
        try:
            doc = json.loads(candidate)
            break
        except ValueError:
            continue
    if not isinstance(doc, dict):
        return empty
    data = doc.get("data") if isinstance(doc.get("data"), dict) else {}
    inner = data.get("state") if isinstance(data.get("state"), str) else None
    if inner is None and isinstance(data.get("status"), str):
        # create_task uses data.status="submitted"; task status uses data.state="running"
        if data["status"].lower() not in {"ok", "submitted", "duplicate_submit_converged"}:
            inner = data["status"]
    return {
        "status": doc.get("status") if isinstance(doc.get("status"), str) else None,
        "project_id": _cli_id(data.get("project_id")),
        "run_id": _cli_id(data.get("run_id")),
        "state": inner.lower() if inner else None,
    }


def _call_deadline(job_deadline: float) -> float:
    return min(job_deadline, time.monotonic() + _CALL_TIMEOUT_SECONDS)


# ---------------------------------------------------------------------------------------
# the detached runner: `python -m app.providers.genspark <job_dir>`
def _kill_tree(proc: subprocess.Popen) -> None:
    try:
        if os.name == "nt":  # the CLI is a .cmd: kill cmd.exe AND the node process below it
            subprocess.run(
                ["taskkill", "/F", "/T", "/PID", str(proc.pid)],
                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=30,
            )
        else:
            import signal

            os.killpg(os.getpgid(proc.pid), signal.SIGKILL)
    except Exception:
        proc.kill()


def _run_logged(command: list[str], cwd: Path, step: str, deadline: float) -> tuple[int | None, bool]:
    """Run `command` with its output in `<step>.stdout.txt` / `<step>.stderr.txt`.
    Returns (exit code or None, timed_out)."""
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0) if os.name == "nt" else 0
    extra = {"start_new_session": True} if os.name != "nt" else {}
    with open(cwd / f"{step}.stdout.txt", "wb") as out, open(cwd / f"{step}.stderr.txt", "wb") as err:
        try:
            proc = subprocess.Popen(
                command, cwd=str(cwd), stdin=subprocess.DEVNULL, stdout=out, stderr=err,
                creationflags=flags, **extra,
            )
        except OSError as exc:
            err.write(f"could not start the command: {exc.__class__.__name__}".encode())
            return -1, False
        while True:
            try:
                return proc.wait(timeout=1), False
            except subprocess.TimeoutExpired:
                if time.monotonic() > deadline:
                    _kill_tree(proc)
                    return None, True


def run_job(job_dir: Path) -> None:
    """Submit a slides task, poll until it finishes, export the PPTX. Never uses --follow."""
    job = json.loads((job_dir / "job.json").read_text(encoding="utf-8"))
    poll_seconds = float(job.get("poll_seconds") or POLL_SECONDS)
    stop = threading.Event()

    def beat() -> None:
        while not stop.is_set():
            try:
                _write_text(job_dir / "heartbeat.txt", str(time.time()))
            except OSError:
                pass
            stop.wait(_HEARTBEAT_SECONDS)

    threading.Thread(target=beat, daemon=True).start()
    deadline = time.monotonic() + float(job["timeout_seconds"])
    report: dict = {
        "create_exit": None, "timed_out": False, "export_attempted": False, "export_exit": None,
        "reported_status": None, "project_id": None, "run_id": None, "task_state": None, "polls": 0,
    }
    try:
        code, timed_out = _run_logged(job["create_command"], job_dir, "create", _call_deadline(deadline))
        report["create_exit"], report["timed_out"] = code, timed_out
        parsed = _parse_cli_json(job_dir / "create.stdout.txt")
        report["reported_status"] = parsed["status"]
        report["project_id"] = parsed["project_id"]
        report["run_id"] = parsed["run_id"]
        if timed_out or parsed["status"] == "error" or (job_dir / OUTPUT_NAME).is_file():
            return
        watch = parsed["run_id"] or parsed["project_id"]
        project_id = parsed["project_id"]
        if not watch:
            return
        failures = 0
        while time.monotonic() < deadline:
            if (job_dir / OUTPUT_NAME).is_file():
                return
            report["polls"] += 1
            step = f"status{report['polls']}"
            code, timed_out = _run_logged(
                job["status_command"] + [watch], job_dir, step, _call_deadline(deadline),
            )
            report["timed_out"] = timed_out
            if timed_out:
                return
            parsed = _parse_cli_json(job_dir / f"{step}.stdout.txt")
            if parsed["project_id"]:
                project_id = parsed["project_id"]
                report["project_id"] = project_id
            if parsed["run_id"]:
                watch = parsed["run_id"]
                report["run_id"] = parsed["run_id"]
            if parsed["state"]:
                report["task_state"] = parsed["state"]
            if parsed["state"] in _FAILED_STATES:
                return
            if parsed["status"] == "error" or (code not in (0, None) and parsed["state"] is None):
                failures += 1
                if failures >= _POLL_MAX_FAILURES:
                    report["reported_status"] = parsed["status"] or "error"
                    return
            else:
                failures = 0
            if parsed["state"] in _ACTIVE_STATES or parsed["state"] is None:
                remaining = deadline - time.monotonic()
                if remaining <= 0:
                    report["timed_out"] = True
                    return
                stop.wait(min(poll_seconds, remaining))
                continue
            export_id = project_id or watch
            report["export_attempted"] = True
            code2, timed_out2 = _run_logged(
                job["export_command"] + [export_id, "--format", "pptx", "-o", OUTPUT_NAME],
                job_dir, "export", _call_deadline(deadline),
            )
            report["export_exit"], report["timed_out"] = code2, timed_out2
            return
        report["timed_out"] = True
    finally:
        stop.set()
        _write_text(job_dir / "result.json", json.dumps(report))  # written last: its existence means "done"


if __name__ == "__main__":  # pragma: no cover - exercised through a real subprocess in the tests
    if len(sys.argv) != 2:
        raise SystemExit("usage: python -m app.providers.genspark <job_dir>")
    run_job(Path(sys.argv[1]))
