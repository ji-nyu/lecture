"""MockPresentationProvider: lets the whole flow run without any external PPT service.

It "renders" a plain, real .pptx (python-pptx): one slide per slide of the approved structure, in
order, with the planned title. It is deliberately not a designed presentation and every slide says
so (`MOCK`): no layout, no images, no lecture content. It never reads the prompt text beyond
hashing/counting what the request already states, so it cannot invent anything.

For tests and demos the job can take several polls (`polls_until_done`) and can fail on purpose
(`fail_at`). Job state lives in `<root>/<job_id>/` so it survives a server restart.
"""

from __future__ import annotations

import json
import logging
import os
import re
import uuid
from pathlib import Path

from ..errors import PresentationGenerationFailed, PresentationProviderUnavailable
from ..models.presentation import (
    DownloadedFile,
    JobState,
    PresentationRequest,
    ProviderResult,
    ProviderStatus,
)
from .base import PresentationProvider, ProviderAvailability

logger = logging.getLogger("ailecturegen")

_JOB_RE = re.compile(r"^mock-[0-9a-f]{32}$")  # also blocks path traversal
PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
FILE_NAME = "mock_presentation.pptx"
MOCK_NOTE = "MOCK provider: placeholder slides with the planned titles only - not a designed presentation."

_STAGES = {
    JobState.queued: "대기 중 (mock)",
    JobState.running: "슬라이드 생성 중 (mock)",
    JobState.completed: "완료 (mock)",
    JobState.failed: "실패 (mock)",
}


class MockPresentationProvider(PresentationProvider):
    name = "mock"
    is_mock = True

    def __init__(self, root: Path, *, polls_until_done: int = 0, fail_at: str | None = None):
        """`polls_until_done`: how many get_status calls until the job ends (0 = already done when
        created). `fail_at`: None | "create" (create_presentation raises) | "render" (the job ends failed)."""
        if fail_at not in (None, "create", "render"):
            raise ValueError("fail_at must be None, 'create' or 'render'")
        self.root = Path(root)
        self.polls_until_done = max(0, int(polls_until_done))
        self.fail_at = fail_at

    # ---- interface --------------------------------------------------------------------
    def availability(self) -> ProviderAvailability:
        return ProviderAvailability(configured=True, message=MOCK_NOTE)

    def create_presentation(self, request: PresentationRequest) -> ProviderStatus:
        if self.fail_at == "create":
            raise PresentationProviderUnavailable("Mock 공급자가 요청을 거절했습니다(테스트용 실패).")
        job_id = f"mock-{uuid.uuid4().hex}"
        job_dir = self._job_dir(job_id)
        job_dir.mkdir(parents=True, exist_ok=True)
        try:
            data = self._render(request)
            (job_dir / FILE_NAME).write_bytes(data)
        except Exception as exc:
            logger.exception("Mock rendering failed", exc_info=exc)  # server log only
            raise PresentationGenerationFailed("Mock 공급자가 슬라이드 파일을 만들지 못했습니다.") from exc
        job = {
            "job_id": job_id,
            "slide_count": request.slide_count,
            "prompt_hash": request.prompt_hash,
            "polls": 0,
            "total_polls": self.polls_until_done,
            "will_fail": self.fail_at == "render",
        }
        self._save(job)
        return self._status(job)

    def get_status(self, job_id: str) -> ProviderStatus:
        job = self._load(job_id)
        job["polls"] += 1  # every poll moves a slow mock job forward
        self._save(job)
        return self._status(job)

    def get_result(self, job_id: str) -> ProviderResult:
        job = self._load(job_id)
        self._require_completed(job)
        data = (self._job_dir(job_id) / FILE_NAME).read_bytes()
        return ProviderResult(
            job_id=job_id, file_name=FILE_NAME, content_type=PPTX_TYPE,
            size_bytes=len(data), slide_count=job["slide_count"],
        )

    def download(self, job_id: str) -> DownloadedFile:
        job = self._load(job_id)
        self._require_completed(job)
        data = (self._job_dir(job_id) / FILE_NAME).read_bytes()
        return DownloadedFile(file_name=FILE_NAME, content_type=PPTX_TYPE, data=data)

    # ---- internals --------------------------------------------------------------------
    def _job_dir(self, job_id: str) -> Path:
        if not _JOB_RE.match(job_id or ""):
            raise PresentationGenerationFailed("알 수 없는 프레젠테이션 작업입니다.")
        return self.root / job_id

    def _load(self, job_id: str) -> dict:
        path = self._job_dir(job_id) / "job.json"
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            raise PresentationGenerationFailed("프레젠테이션 작업 정보를 찾을 수 없습니다.") from exc

    def _save(self, job: dict) -> None:
        path = self._job_dir(job["job_id"]) / "job.json"
        tmp = path.with_suffix(".tmp")
        tmp.write_text(json.dumps(job), encoding="utf-8")
        os.replace(tmp, path)

    @staticmethod
    def _state(job: dict) -> tuple[JobState, int]:
        total, polls = job["total_polls"], job["polls"]
        if polls >= total:
            return (JobState.failed if job["will_fail"] else JobState.completed), (0 if job["will_fail"] else 100)
        return (JobState.queued if polls == 0 else JobState.running), int(100 * polls / total)

    def _status(self, job: dict) -> ProviderStatus:
        state, progress = self._state(job)
        return ProviderStatus(
            job_id=job["job_id"], state=state, progress=progress, stage=_STAGES[state],
            message="Mock 공급자가 슬라이드를 만들지 못했습니다(테스트용 실패)." if state == JobState.failed else None,
        )

    def _require_completed(self, job: dict) -> None:
        if self._state(job)[0] != JobState.completed:
            raise PresentationGenerationFailed("프레젠테이션이 아직 완성되지 않았습니다.")

    @staticmethod
    def _render(request: PresentationRequest) -> bytes:
        import io

        from pptx import Presentation
        from pptx.util import Inches, Pt

        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        prs.core_properties.title = request.title
        prs.core_properties.author = "AILectureGen (mock provider)"
        prs.core_properties.comments = MOCK_NOTE
        layout = prs.slide_layouts[5]  # "Title Only"
        total = len(request.slides)
        for brief in request.slides:
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = brief.title
            info = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(11.5), Inches(1.0))
            info.text_frame.text = f"슬라이드 {brief.slide_number}/{total} · {brief.slide_type}"
            info.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
            foot = slide.shapes.add_textbox(Inches(0.8), Inches(6.6), Inches(11.5), Inches(0.5))
            foot.text_frame.text = "MOCK PROVIDER - 자리표시 슬라이드입니다. 실제 디자인된 프레젠테이션이 아닙니다."
            foot.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()
