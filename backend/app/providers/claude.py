"""ClaudeProvider (STAGE 8): renders the presentation with Claude (Anthropic API + the `pptx` Agent Skill).

All Claude-specific code lives in this file and nowhere else; the Lecture Engine only knows
`PresentationProvider`. (This replaced the Genspark CLI provider.)

What was verified before writing this (MASTER_SPEC STAGE 8: no guessed endpoints, parameters or
response schemas). Sources: the official docs "Agent Skills quickstart (PowerPoint)" and "Using Agent
Skills with the API" on platform.claude.com, the "Models overview", and the installed `anthropic`
SDK (1.x: `client.messages.create(container=...)`, `client.files.retrieve_metadata/download`).
  * Request: Messages API, `container={"skills": [{"type": "anthropic", "skill_id": "pptx",
    "version": "latest"}]}` plus the code execution tool (`{"type": "code_execution_20250825",
    "name": "code_execution"}`; the quickstart says any current code execution tool version satisfies
    the Skills requirement). No beta header is used: the current docs' examples send none.
  * Long turns: the API may stop with `stop_reason == "pause_turn"`. The documented way on is to append
    the assistant content to `messages` and call again with `container={"id": response.container.id,
    "skills": [...]}` until the stop reason is something else.
  * Result: generated files appear as `bash_code_execution_tool_result` blocks whose `content.type` is
    `bash_code_execution_result` and whose `content.content[]` items carry a `file_id`. The file is
    fetched with the Files API (metadata for the file name, then `download`).
  * Limits that matter: the Skills sandbox has no network access (so no stock photos or downloads:
    the prompt already tells the renderer to draw visuals itself), and the model must support code
    execution (CLAUDE_MODEL).
  * NOT verified: no live call was ever made (no key was available while writing this). The tests run
    the real request/response code path against a local fake Anthropic server that returns the
    documented shapes, never against Anthropic.

How a job runs (a generation takes minutes, the API must not block):
  create_presentation writes the parameters to `<jobs>/<job_id>/args.json` and starts a detached runner
  (`python -m app.providers.claude <job_dir>`) that calls the API, downloads the file to `output.pptx`
  and finally writes `result.json`. get_status only reads those files, so a job survives a server
  restart. A heartbeat file tells a dead runner from a slow one. The API reports no progress, so
  progress stays 0.

Security: the API key is handed to the runner through its environment only (never a command line, never
a file). Nothing returned to a client contains API error text; error details go to the server log,
redacted.
"""

from __future__ import annotations

import importlib.util
import json
import logging
import os
import re
import subprocess
import sys
import threading
import time
import uuid
from pathlib import Path
from typing import Callable

from ..errors import PresentationGenerationFailed, PresentationProviderUnavailable
from ..models.presentation import (
    DownloadedFile,
    JobState,
    PresentationRequest,
    ProviderResult,
    ProviderStatus,
)
from .base import PresentationProvider, ProviderAvailability

logger = logging.getLogger("ailecturegen")

PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
OUTPUT_NAME = "output.pptx"
FILE_NAME = "claude_presentation.pptx"
SKILL = {"type": "anthropic", "skill_id": "pptx", "version": "latest"}
CODE_EXECUTION_TOOL = {"type": "code_execution_20250825", "name": "code_execution"}
MAX_TURNS = 25  # 1 request + up to 24 `pause_turn` continuations

_BACKEND_DIR = Path(__file__).resolve().parents[2]  # the directory that contains the `app` package
_JOB_RE = re.compile(r"^cld-[0-9a-f]{32}$")  # also blocks path traversal
_URL_RE = re.compile(r"^https?://[^\s]+$")
_HEARTBEAT_SECONDS = 5
_HEARTBEAT_STALE_SECONDS = 60  # no heartbeat and no result for this long: the runner is gone
_NOT_FOR_RUNNER = ("LLM_API_KEY", "OPENAI_API_KEY", "ANTHROPIC_AUTH_TOKEN")  # the runner gets ANTHROPIC_API_KEY only

MSG_NO_KEY = "ANTHROPIC_API_KEY가 설정되지 않았습니다. backend/.env에 Anthropic API 키를 넣어 주세요."
MSG_NO_SDK = "anthropic 패키지가 설치되어 있지 않습니다. `pip install -r requirements.txt` 를 실행해 주세요."
MSG_NO_MODEL = "CLAUDE_MODEL이 비어 있습니다. 코드 실행을 지원하는 Claude 모델 ID를 설정해 주세요."
MSG_BAD_URL = "CLAUDE_BASE_URL은 http:// 또는 https:// 로 시작하는 주소여야 합니다."
HINT = "서버 로그를 확인해 주세요."

# What the runner records (`error_kind`) -> what a client may read. No API text is ever passed on.
_KIND_MESSAGES = {
    "auth": "Anthropic이 API 키를 거부했습니다. backend/.env의 ANTHROPIC_API_KEY를 확인해 주세요.",
    "permission": "이 API 키로는 요청한 모델이나 기능(Agent Skills, 코드 실행)을 쓸 수 없습니다. 키 권한과 CLAUDE_MODEL을 확인해 주세요.",
    "rate_limit": "Anthropic 요청 한도에 걸렸습니다. 잠시 후 다시 생성해 주세요.",
    "connection": "Anthropic API에 연결하지 못했습니다. 네트워크를 확인하고 다시 생성해 주세요.",
    "not_found": "모델 또는 Skill을 찾을 수 없습니다. CLAUDE_MODEL을 확인해 주세요.",
    "bad_request": f"Anthropic이 요청을 거부했습니다(크레딧 잔액, 모델, 요청 크기를 확인해 주세요). {HINT}",
    "server": "Anthropic 서버가 일시적으로 응답하지 않습니다. 잠시 후 다시 생성해 주세요.",
    "max_tokens": "출력 한도(CLAUDE_MAX_TOKENS)에 도달해 PPTX가 완성되지 않았습니다. 값을 늘리거나 슬라이드 수를 줄여 주세요.",
    "refusal": f"Claude가 요청을 처리하지 않았습니다. {HINT}",
    "too_many_turns": f"Claude가 정해진 횟수 안에 PPTX를 끝내지 못했습니다. {HINT}",
    "no_file": f"Claude가 PPTX 파일을 만들지 못했습니다. {HINT}",
    "no_pptx": f"Claude가 만든 파일 중 PPTX가 없습니다. {HINT}",
    "unexpected": f"프레젠테이션 생성 중 예기치 못한 오류가 발생했습니다. {HINT}",
}


def build_instructions(request: PresentationRequest) -> str:
    """The `system` prompt (working rules): only what the spec gives the renderer to do and not to do,
    plus the fixed slide list. The lecture itself is the user message."""
    titles = "\n".join(f"{s.slide_number}. {s.title}" for s in request.slides)
    return (
        "사용자 메시지는 이미 확정된 강의 명세입니다. 당신은 그것을 PowerPoint(.pptx)로 렌더링하는 역할만 맡습니다. "
        "pptx Skill로 파일을 정확히 한 개 만드세요.\n"
        f"- 슬라이드는 정확히 {request.slide_count}장이며, 아래 목록의 순서와 제목을 그대로 따릅니다. "
        "슬라이드를 추가, 삭제, 병합, 분할하지 않습니다.\n"
        "- 맡는 일: 시각 레이아웃, 슬라이드 렌더링, 프레젠테이션 스타일링.\n"
        "- 맡지 않는 일: 강의 구조, 중요 개념, 섹션 시간, 난이도를 정하거나 바꾸는 것. "
        "명세에 적힌 텍스트와 수치를 그대로 사용하고 새로운 사실이나 예시를 덧붙이지 않습니다.\n"
        "- 각 슬라이드의 본문 항목은 화면에 읽히는 글자로 넣습니다. "
        "제목이나 개념 이름만 적힌 빈 상자·빈 카드를 만들지 않으며, 장식 이미지가 본문을 대체하지 않습니다.\n"
        "- 실행 환경에는 인터넷이 없습니다. 사진이나 외부 이미지를 가져오지 말고, 도형·표·차트·다이어그램으로 직접 그립니다.\n"
        "- 결과는 편집 가능한 .pptx 한 개여야 합니다(이미지로 만든 슬라이드 금지). 끝나면 한두 문장으로만 답합니다.\n\n"
        f"슬라이드 목록:\n{titles}"
    )


def build_prompt(request: PresentationRequest) -> str:
    """The user message: the whole prompt verbatim plus a closing sentence that states the deliverable."""
    return (
        f"{request.final_prompt}\n\n---\n"
        f"위 명세에 따라 정확히 {request.slide_count}장의 프레젠테이션(.pptx)을 만들어 주세요."
    )


class ClaudeProvider(PresentationProvider):
    name = "claude"
    is_mock = False

    def __init__(
        self,
        jobs_root: Path,
        *,
        api_key: str | None = None,
        model: str = "claude-sonnet-5-5",
        max_tokens: int = 16000,
        base_url: str | None = None,
        timeout_seconds: int = 1800,
        launcher: Callable[[Path], None] | None = None,
    ):
        self.root = Path(jobs_root)
        self._api_key = api_key or None  # never logged, never written to disk
        self.model = (model or "").strip()
        self.max_tokens = int(max_tokens)
        self.base_url = base_url or None
        self.timeout_seconds = int(timeout_seconds)
        self._launcher = launcher or self._launch_runner  # tests run the runner in a thread instead
        self._logged: set[str] = set()

    def __repr__(self) -> str:  # the key must never show up in a log line
        return f"ClaudeProvider(model={self.model!r}, key={'set' if self._api_key else 'unset'})"

    # ---- interface --------------------------------------------------------------------
    def availability(self) -> ProviderAvailability:
        if not self._api_key:
            return ProviderAvailability(configured=False, message=MSG_NO_KEY)
        if not self.model:
            return ProviderAvailability(configured=False, message=MSG_NO_MODEL)
        if self.base_url and not _URL_RE.match(self.base_url):
            return ProviderAvailability(configured=False, message=MSG_BAD_URL)
        if importlib.util.find_spec("anthropic") is None:
            return ProviderAvailability(configured=False, message=MSG_NO_SDK)
        return ProviderAvailability(configured=True)

    def create_presentation(self, request: PresentationRequest) -> ProviderStatus:
        availability = self.availability()
        if not availability.configured:
            raise PresentationProviderUnavailable(availability.message)
        job_id = f"cld-{uuid.uuid4().hex}"
        job_dir = self.root / job_id
        try:
            job_dir.mkdir(parents=True, exist_ok=True)
            _write_text(job_dir / "args.json", json.dumps(
                {
                    "model": self.model,
                    "max_tokens": self.max_tokens,
                    "system": build_instructions(request),
                    "prompt": build_prompt(request),
                },
                ensure_ascii=False,
            ))
            _write_text(job_dir / "job.json", json.dumps(
                {
                    "job_id": job_id,
                    "created_at": time.time(),
                    "slide_count": request.slide_count,
                    "prompt_hash": request.prompt_hash,
                    "timeout_seconds": self.timeout_seconds,
                    # No key in here: the runner gets it through its environment.
                }
            ))
            self._launcher(job_dir)
        except Exception as exc:
            logger.exception("Could not start the Claude runner", exc_info=exc)  # server log only
            raise PresentationProviderUnavailable(f"Claude 작업을 시작하지 못했습니다. {HINT}") from exc
        return ProviderStatus(job_id=job_id, state=JobState.queued, stage="Claude 작업 시작")

    def get_status(self, job_id: str) -> ProviderStatus:
        job_dir, job = self._job(job_id)
        state, message, _ = self._evaluate(job_id, job_dir, job)
        if state == JobState.completed:
            return ProviderStatus(job_id=job_id, state=state, progress=100, stage="완료")
        if state == JobState.failed:
            return ProviderStatus(job_id=job_id, state=state, stage="실패", message=message)
        started = (job_dir / "heartbeat.txt").is_file()
        return ProviderStatus(
            job_id=job_id,
            state=JobState.running if started else JobState.queued,
            stage="Claude가 슬라이드를 생성 중 (진행률은 제공되지 않음)" if started else "Claude 작업 시작",
        )

    def get_result(self, job_id: str) -> ProviderResult:
        job_dir, job = self._job(job_id)
        state, message, slides = self._evaluate(job_id, job_dir, job)
        if state != JobState.completed:
            raise PresentationGenerationFailed(message or "프레젠테이션이 아직 완성되지 않았습니다.")
        return ProviderResult(
            job_id=job_id, file_name=FILE_NAME, content_type=PPTX_TYPE,
            size_bytes=(job_dir / OUTPUT_NAME).stat().st_size, slide_count=slides,
        )

    def download(self, job_id: str) -> DownloadedFile:
        job_dir, job = self._job(job_id)
        state, message, _ = self._evaluate(job_id, job_dir, job)
        if state != JobState.completed:
            raise PresentationGenerationFailed(message or "프레젠테이션이 아직 완성되지 않았습니다.")
        return DownloadedFile(
            file_name=FILE_NAME, content_type=PPTX_TYPE, data=(job_dir / OUTPUT_NAME).read_bytes()
        )

    # ---- internals --------------------------------------------------------------------
    def _job(self, job_id: str) -> tuple[Path, dict]:
        if not _JOB_RE.match(job_id or ""):
            raise PresentationGenerationFailed("알 수 없는 프레젠테이션 작업입니다.")
        job_dir = self.root / job_id
        try:
            return job_dir, json.loads((job_dir / "job.json").read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            raise PresentationGenerationFailed("프레젠테이션 작업 정보를 찾을 수 없습니다.") from exc

    def _child_env(self) -> dict[str, str]:
        env = {k: v for k, v in os.environ.items() if k not in _NOT_FOR_RUNNER}
        env["ANTHROPIC_API_KEY"] = self._api_key or ""  # the SDK reads this variable by itself
        if self.base_url:
            env["ANTHROPIC_BASE_URL"] = self.base_url
        else:
            env.pop("ANTHROPIC_BASE_URL", None)  # an unrelated proxy setting must not redirect the key
        return env

    def _launch_runner(self, job_dir: Path) -> None:
        kwargs: dict = dict(
            cwd=str(_BACKEND_DIR), env=self._child_env(), close_fds=True,
            stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        )
        if os.name == "nt":  # outlive this process, no console window
            kwargs["creationflags"] = subprocess.DETACHED_PROCESS | subprocess.CREATE_NEW_PROCESS_GROUP
        else:
            kwargs["start_new_session"] = True
        subprocess.Popen([sys.executable, "-m", "app.providers.claude", str(job_dir)], **kwargs)

    def _redact(self, text: str) -> str:
        return text.replace(self._api_key, "***") if self._api_key else text

    def _log_failure(self, job_id: str, job_dir: Path, result: dict) -> None:
        if job_id in self._logged:  # status is polled: one log line per job is enough
            return
        self._logged.add(job_id)
        try:
            detail = (job_dir / "error.txt").read_text(encoding="utf-8", errors="replace")[-800:].strip()
        except OSError:
            detail = ""
        logger.warning(
            "Claude job %s failed (%s; stop_reason=%s): %s",
            job_id, result.get("error_kind"), result.get("stop_reason"), self._redact(detail),
        )

    def _evaluate(self, job_id: str, job_dir: Path, job: dict) -> tuple[JobState, str | None, int | None]:
        """(state, user-facing message if failed, slide count of the finished file)."""
        result = _read_json(job_dir / "result.json")
        if result is None:  # still running - unless the runner died
            beat = job_dir / "heartbeat.txt"
            try:
                last = beat.stat().st_mtime if beat.is_file() else float(job["created_at"])
            except (OSError, ValueError, KeyError):
                last = 0.0
            if time.time() - last > _HEARTBEAT_STALE_SECONDS:
                return JobState.failed, "Claude 작업을 처리하던 프로세스가 중단되었습니다. 다시 생성해 주세요.", None
            return JobState.running, None, None

        if result.get("timed_out"):
            minutes = max(1, int(job.get("timeout_seconds", self.timeout_seconds)) // 60)
            return JobState.failed, f"Claude 생성이 제한 시간({minutes}분) 안에 끝나지 않아 중단했습니다.", None
        if result.get("outcome") != "done":
            self._log_failure(job_id, job_dir, result)
            kind = result.get("error_kind")
            return JobState.failed, _KIND_MESSAGES.get(kind, _KIND_MESSAGES["unexpected"]), None
        output = job_dir / OUTPUT_NAME
        slides = _count_slides(output) if output.is_file() else None
        if slides is None:
            self._log_failure(job_id, job_dir, {"error_kind": "invalid_pptx", "stop_reason": result.get("stop_reason")})
            return JobState.failed, "Claude가 돌려준 파일이 올바른 PPTX가 아닙니다.", None
        return JobState.completed, None, slides


# ---------------------------------------------------------------------------------------
# helpers shared with the runner
def _write_text(path: Path, text: str) -> None:
    tmp = path.with_name(path.name + ".tmp")
    tmp.write_text(text, encoding="utf-8")
    os.replace(tmp, path)


def _read_json(path: Path) -> dict | None:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    return data if isinstance(data, dict) else None


def _count_slides(path: Path) -> int | None:
    """Number of slides if `path` is a readable presentation, else None."""
    try:
        from pptx import Presentation

        return len(Presentation(str(path)).slides)
    except Exception:
        return None


# ---------------------------------------------------------------------------------------
# the detached runner: `python -m app.providers.claude <job_dir>`
def _classify(exc: Exception) -> str:
    import anthropic

    if isinstance(exc, anthropic.AuthenticationError):
        return "auth"
    if isinstance(exc, anthropic.PermissionDeniedError):
        return "permission"
    if isinstance(exc, anthropic.RateLimitError):
        return "rate_limit"
    if isinstance(exc, anthropic.APITimeoutError):
        return "timeout"
    if isinstance(exc, anthropic.APIConnectionError):
        return "connection"
    if isinstance(exc, anthropic.NotFoundError):
        return "not_found"
    if isinstance(exc, anthropic.BadRequestError):
        return "bad_request"
    if isinstance(exc, anthropic.APIStatusError):
        return "server"
    return "unexpected"


def _file_ids(response) -> list[str]:
    """file_ids of the files the code execution tool produced (the documented response shape)."""
    ids: list[str] = []
    for block in getattr(response, "content", None) or []:
        if getattr(block, "type", None) != "bash_code_execution_tool_result":
            continue
        inner = getattr(block, "content", None)
        if getattr(inner, "type", None) != "bash_code_execution_result":
            continue
        for item in getattr(inner, "content", None) or []:
            file_id = getattr(item, "file_id", None)
            if file_id:
                ids.append(file_id)
    return ids


def _generate(client, args: dict, job_dir: Path, deadline: float) -> dict:
    """One generation: request, `pause_turn` continuations, file download. Returns the result report."""
    messages: list[dict] = [{"role": "user", "content": args["prompt"]}]
    container: dict = {"skills": [SKILL]}
    file_ids: list[str] = []
    report: dict = {"outcome": "error", "turns": 0, "input_tokens": 0, "output_tokens": 0}
    response = None
    for _ in range(MAX_TURNS):
        remaining = deadline - time.monotonic()
        if remaining <= 1:
            return {**report, "timed_out": True}
        response = client.messages.create(
            model=args["model"],
            max_tokens=int(args["max_tokens"]),
            system=args["system"],
            container=container,
            messages=messages,
            tools=[CODE_EXECUTION_TOOL],
            timeout=remaining,  # an explicit timeout also lifts the SDK's "use streaming" guard
        )
        report["turns"] += 1
        usage = getattr(response, "usage", None)
        report["input_tokens"] += int(getattr(usage, "input_tokens", 0) or 0)
        report["output_tokens"] += int(getattr(usage, "output_tokens", 0) or 0)
        file_ids += _file_ids(response)
        report["stop_reason"] = getattr(response, "stop_reason", None)
        if response.stop_reason != "pause_turn":
            break
        messages.append({"role": "assistant", "content": response.content})
        container = {"id": response.container.id, "skills": [SKILL]}
    else:
        return {**report, "error_kind": "too_many_turns"}

    if not file_ids:
        stop = report.get("stop_reason")
        kind = "max_tokens" if stop == "max_tokens" else "refusal" if stop == "refusal" else "no_file"
        return {**report, "error_kind": kind}

    for file_id in reversed(file_ids):  # the newest file first
        if time.monotonic() > deadline:
            return {**report, "timed_out": True}
        meta = client.files.retrieve_metadata(file_id=file_id)
        if not str(getattr(meta, "filename", "") or "").lower().endswith(".pptx"):
            continue
        part = job_dir / (OUTPUT_NAME + ".part")
        client.files.download(file_id=file_id).write_to_file(part)
        os.replace(part, job_dir / OUTPUT_NAME)
        return {**report, "outcome": "done"}
    return {**report, "error_kind": "no_pptx"}


def run_job(job_dir: Path, client=None) -> None:
    """Runs one job to its end and writes `result.json` (last: its existence means "done").
    `client` is injectable for tests; the real runner builds one from the environment."""
    job = json.loads((job_dir / "job.json").read_text(encoding="utf-8"))
    args = json.loads((job_dir / "args.json").read_text(encoding="utf-8"))
    stop = threading.Event()

    def beat() -> None:
        while not stop.is_set():
            try:
                _write_text(job_dir / "heartbeat.txt", str(time.time()))
            except OSError:
                pass
            stop.wait(_HEARTBEAT_SECONDS)

    threading.Thread(target=beat, daemon=True).start()
    deadline = time.monotonic() + float(job["timeout_seconds"])
    report: dict = {"outcome": "error", "error_kind": "unexpected"}
    try:
        if client is None:
            import anthropic

            # Reads ANTHROPIC_API_KEY (and ANTHROPIC_BASE_URL) from the environment, nowhere else.
            client = anthropic.Anthropic(max_retries=2)
        report = _generate(client, args, job_dir, deadline)
    except Exception as exc:
        try:
            kind = _classify(exc)
        except Exception:
            kind = "unexpected"
        report = {"outcome": "error", "error_kind": kind, "timed_out": kind == "timeout"}
        detail = f"{exc.__class__.__name__} status={getattr(exc, 'status_code', None)}: {str(exc)[:500]}"
        key = os.environ.get("ANTHROPIC_API_KEY")
        if key:
            detail = detail.replace(key, "***")
        try:
            _write_text(job_dir / "error.txt", detail)
        except OSError:
            pass
    finally:
        stop.set()
        _write_text(job_dir / "result.json", json.dumps(report))


if __name__ == "__main__":  # pragma: no cover - exercised through a real subprocess in the tests
    if len(sys.argv) != 2:
        raise SystemExit("usage: python -m app.providers.claude <job_dir>")
    run_job(Path(sys.argv[1]))
