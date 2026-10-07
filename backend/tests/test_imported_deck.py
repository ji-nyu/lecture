"""Instructor-made PPTX: scripts and video use that deck as-is."""

from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from app.config import Settings
from app.main import create_app


def _deck(tmp_path: Path) -> Path:
    prs = Presentation()
    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    one = prs.slides.add_slide(layout)
    box = one.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(8), Inches(1))
    box.text_frame.text = "AIoT 입문"
    chrome = one.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(8), Inches(2))
    chrome.text_frame.text = "20분짜리 강의 — AIoT → MQTT — 01 / 02\n오늘은 AIoT의 기초 흐름을 안내합니다."
    two = prs.slides.add_slide(layout)
    title = two.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(8), Inches(1))
    title.text_frame.text = "용어 정의: AIoT"
    body = two.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(8), Inches(3))
    body.text_frame.text = (
        "AIoT는 Artificial Intelligence of Things의 약자로, "
        "사물인터넷에 인공지능을 결합한 체계를 말한다."
    )
    path = tmp_path / "lecture.pptx"
    prs.save(path)
    return path


def _deck_with_repeated_definition(tmp_path: Path) -> Path:
    prs = Presentation()
    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    one = prs.slides.add_slide(layout)
    one.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(8), Inches(1)).text_frame.text = "용어 정의: AIoT"
    one.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(8), Inches(2)).text_frame.text = (
        "AIoT는 Artificial Intelligence of Things의 약자로, "
        "사물인터넷에 인공지능을 결합한 체계를 말한다."
    )
    two = prs.slides.add_slide(layout)
    two.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(8), Inches(1)).text_frame.text = "AIoT 구성"
    two.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(8), Inches(3)).text_frame.text = (
        "AIoT는 Artificial Intelligence of Things의 약자로, "
        "사물인터넷에 인공지능을 결합한 체계를 말한다.\n"
        "센서는 물리량을 측정하고 서버는 저장과 후처리를 담당한다."
    )
    path = tmp_path / "repeat.pptx"
    prs.save(path)
    return path


@pytest.fixture()
def client(tmp_path):
    return TestClient(create_app(Settings(data_dir=tmp_path, max_upload_mb=5, cors_origins=())), raise_server_exceptions=False)


def test_imported_deck_makes_scripts_and_keeps_the_pptx(client, tmp_path):
    deck = _deck(tmp_path)
    pid = client.post("/projects", json={"title": "가져온 강의"}).json()["id"]
    with deck.open("rb") as fh:
        r = client.post(
            f"/projects/{pid}/deck",
            files={"file": ("lecture.pptx", fh, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
            data={"duration_minutes": "20"},
        )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["input_mode"] == "deck"
    assert body["has_plan"] and body["has_slides"] and body["has_presentation"]
    assert body["presentation_status"] == "completed"
    spec = client.get(f"/projects/{pid}/slides").json()
    assert spec["slide_count"] == 2
    assert spec["slides"][1]["title"] == "용어 정의: AIoT"
    assert spec["slides"][1]["slide_type"] == "definition"
    assert 1 <= sum(s["estimated_explanation_time"] for s in spec["slides"]) <= 1200
    scripts = client.get(f"/projects/{pid}/scripts")
    assert scripts.status_code == 200, scripts.text
    spoken = scripts.json()
    assert spoken["slide_count"] == 2
    first, second = spoken["slides"]
    assert "기초 흐름" in first["text"]
    assert "Artificial Intelligence" not in first["text"]
    assert "사물인터넷" in second["text"] or "Artificial Intelligence" in second["text"]
    assert "기초 흐름" not in second["text"]
    for s in spoken["slides"]:
        assert "(" not in s["text"]
        assert "20분" not in s["text"]
        assert "분짜리" not in s["text"]
        for mark in ("—", "→", "/", "─", "·"):
            assert mark not in s["text"], (s["slide_number"], mark, s["text"])
        assert "다시 말하면, 첫째 사례" not in s["text"]
        assert "다시 말하면, 사례" not in s["text"]
        assert "학습 비교입니다" not in s["text"]
        assert "만들어진…입니다" not in s["text"]
        assert "빠르지만입니다" not in s["text"]
        assert "유실 가능입니다" not in s["text"]
        gap = abs(s["spoken_seconds"] - s["estimated_seconds"])
        assert gap <= max(15, int(s["estimated_seconds"] * 0.35)), (
            s["slide_number"],
            s["spoken_seconds"],
            s["estimated_seconds"],
        )
    dl = client.get(f"/projects/{pid}/presentation/download")
    assert dl.status_code == 200
    assert dl.content[:4] == b"PK\x03\x04"


def test_imported_scripts_do_not_repeat_the_previous_slide(client, tmp_path):
    deck = _deck_with_repeated_definition(tmp_path)
    pid = client.post("/projects", json={"title": "반복 강의"}).json()["id"]
    with deck.open("rb") as fh:
        r = client.post(
            f"/projects/{pid}/deck",
            files={"file": ("repeat.pptx", fh, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
            data={"duration_minutes": "20"},
        )
    assert r.status_code == 200, r.text
    spoken = client.get(f"/projects/{pid}/scripts").json()["slides"]
    assert len(spoken) == 2
    assert "Artificial Intelligence" in spoken[0]["text"] or "사물인터넷" in spoken[0]["text"]
    assert "센서는 물리량" in spoken[1]["text"] or "후처리" in spoken[1]["text"]
    assert "Artificial Intelligence" not in spoken[1]["text"]


def test_deck_upload_rejects_a_markdown_file(client):
    pid = client.post("/projects", json={}).json()["id"]
    r = client.post(
        f"/projects/{pid}/deck",
        files={"file": ("notes.md", b"# hi", "text/markdown")},
        data={"duration_minutes": "20"},
    )
    assert r.status_code == 415
    assert r.json()["error"]["code"] == "DeckNeedsPptx"


def test_changing_duration_on_an_imported_deck_keeps_the_pptx(client, tmp_path):
    deck = _deck(tmp_path)
    pid = client.post("/projects", json={}).json()["id"]
    with deck.open("rb") as fh:
        assert client.post(
            f"/projects/{pid}/deck",
            files={"file": ("lecture.pptx", fh, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
            data={"duration_minutes": "20"},
        ).status_code == 200
    r = client.put(
        f"/projects/{pid}/profile",
        json={
            "audience_level": "general",
            "duration_minutes": 40,
            "difficulty": "introductory",
            "lecture_type": "theory",
            "explanation_depth": "concise",
            "source_policy": "source_first",
        },
    )
    assert r.status_code == 200, r.text
    spec = client.get(f"/projects/{pid}/slides").json()
    assert spec["duration_minutes"] == 40
    assert 1 <= sum(s["estimated_explanation_time"] for s in spec["slides"]) <= 2400
    project = client.get(f"/projects/{pid}").json()
    assert project["has_presentation"] is True
    assert project["input_mode"] == "deck"
    assert client.get(f"/projects/{pid}/presentation/download").status_code == 200


def test_imported_deck_can_choose_an_intro(client, tmp_path):
    deck = _deck(tmp_path)
    pid = client.post("/projects", json={}).json()["id"]
    with deck.open("rb") as fh:
        assert client.post(
            f"/projects/{pid}/deck",
            files={"file": ("lecture.pptx", fh, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
        ).status_code == 200
    assert client.post(
        "/video-intros",
        files={"file": ("open.mp4", b"clip", "video/mp4")},
    ).status_code == 200
    r = client.put(
        f"/projects/{pid}/profile",
        json={
            "audience_level": "general",
            "duration_minutes": 20,
            "difficulty": "introductory",
            "lecture_type": "theory",
            "explanation_depth": "concise",
            "source_policy": "source_first",
            "video_intro": "include",
            "video_intro_file": "open.mp4",
        },
    )
    assert r.status_code == 200, r.text
    assert r.json()["video_intro"] == "include"
    assert r.json()["video_intro_file"] == "open.mp4"
    project = client.get(f"/projects/{pid}").json()
    assert project["input_mode"] == "deck"
    assert project["has_presentation"] is True
