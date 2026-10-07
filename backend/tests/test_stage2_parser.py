"""DocumentParser tests: every supported format is generated as a real file."""

from pathlib import Path

import pytest

from app.errors import DocumentParsingError, EmptyDocument
from app.services.document_parser import DocumentParser

parser = DocumentParser()


def parse(path: Path):
    return parser.parse(path, path.name, "test")


# ---------------------------------------------------------------- txt / md
def test_txt_markdown_headings_hierarchy_and_lines(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text(
        "# 제목\n\n도입 문장이다.\n\n## 1. 첫째\n\n첫째 내용이다.\n\n## 2. 둘째\n\n- 항목 하나\n- 항목 둘\n",
        encoding="utf-8",
    )
    m = parse(f)
    assert m.title == "제목"
    assert [s.title for s in m.sections] == ["제목", "1. 첫째", "2. 둘째"]
    assert [s.level for s in m.sections] == [1, 2, 2]
    assert m.sections[1].parent_id == m.sections[0].id
    assert m.sections[2].blocks[0].kind == "list_item"
    assert m.sections[1].blocks[0].line == 7  # source line numbers are kept
    assert len(m.headings) == 3


def test_plain_numbered_headings_without_markdown(tmp_path):
    f = tmp_path / "plain.txt"
    f.write_text(
        "1. 개요\n\n이것은 개요다.\n\n2. 상세\n\n1. 첫 단계를 수행한다.\n2. 다음 단계를 수행한다.\n",
        encoding="utf-8",
    )
    m = parse(f)
    # numbered *sentences* stay list items; short numbered lines become headings
    assert [h.text for h in m.headings] == ["1. 개요", "2. 상세"]
    assert [b.kind for b in m.sections[1].blocks] == ["list_item", "list_item"]


def test_content_before_first_heading_gets_intro_section(tmp_path):
    f = tmp_path / "x.md"
    f.write_text("머리말 문장이다.\n\n## 본문\n\n내용이다.\n", encoding="utf-8")
    m = parse(f)
    assert m.sections[0].title == "도입"
    assert m.sections[1].title == "본문"


def test_code_tables_formulas_extracted(tmp_path):
    f = tmp_path / "x.md"
    f.write_text(
        "# T\n\n## 코드\n\n```python\nprint('hi')\n```\n\n"
        "| a | b |\n|---|---|\n| 1 | 2 |\n\n"
        "$$E = mc^2$$\n\n인라인 $x^2 + y^2$ 수식이다.\n",
        encoding="utf-8",
    )
    m = parse(f)
    assert m.code_blocks[0].language == "python"
    assert "print('hi')" in m.code_blocks[0].code
    assert m.tables[0].rows == [["a", "b"], ["1", "2"]]
    exprs = [x.expression for x in m.formulas]
    assert "E = mc^2" in exprs and "x^2 + y^2" in exprs
    assert m.code_blocks[0].location.line == 5
    # code is not mixed into the prose of its section
    assert "print" not in m.sections[1].text


def test_cp949_text_is_decoded(tmp_path):
    f = tmp_path / "k.txt"
    f.write_bytes("# 한글 제목\n\n본문이다.\n".encode("cp949"))
    m = parse(f)
    assert m.title == "한글 제목"
    assert any("cp949" in w for w in m.warnings)


def test_empty_or_blank_text_raises_empty_document(tmp_path):
    f = tmp_path / "e.txt"
    f.write_text("  \n\n \n", encoding="utf-8")
    with pytest.raises(EmptyDocument):
        parse(f)


# ------------------------------------------------------------------- docx
def test_docx_headings_lists_tables(tmp_path):
    from docx import Document

    d = Document()
    d.add_heading("Lecture Title", 0)
    d.add_heading("Section A", 1)
    d.add_paragraph("Body text of A.")
    d.add_paragraph("bullet one", style="List Bullet")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "h1", "h2"
    t.cell(1, 0).text, t.cell(1, 1).text = "v1", "v2"
    d.add_heading("Section B", 1)
    d.add_paragraph("Body text of B.")
    f = tmp_path / "a.docx"
    d.save(f)

    m = parse(f)
    assert m.file_type == "docx"
    assert [s.title for s in m.sections] == ["Lecture Title", "Section A", "Section B"]
    assert m.title == "Lecture Title"
    kinds = [b.kind for b in m.sections[1].blocks]
    assert kinds == ["paragraph", "list_item"]
    assert m.tables[0].rows == [["h1", "h2"], ["v1", "v2"]]
    assert m.tables[0].section_id == m.sections[1].id


# ------------------------------------------------------------------- pptx
def test_pptx_slides_notes_tables(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "What is MQTT"
    s1.placeholders[1].text_frame.text = "Lightweight protocol"
    s1.placeholders[1].text_frame.add_paragraph().text = "Publish/Subscribe model"
    s1.notes_slide.notes_text_frame.text = "Explain with a sensor example."
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "QoS levels"
    tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text = "level", "meaning"
    tbl.cell(1, 0).text, tbl.cell(1, 1).text = "0", "at most once"
    f = tmp_path / "a.pptx"
    prs.save(f)

    m = parse(f)
    assert m.file_type == "pptx" and m.page_count == 2
    assert [s.title for s in m.sections] == ["What is MQTT", "QoS levels"]
    assert m.sections[0].page == 1
    texts = [b.text for b in m.sections[0].blocks]
    assert "Lightweight protocol" in texts and "Publish/Subscribe model" in texts
    assert any(b.kind == "note" and "sensor" in b.text for b in m.sections[0].blocks)
    assert m.tables[0].rows[1] == ["0", "at most once"]


# -------------------------------------------------------------------- pdf
def _make_pdf(pages: list[list[str]]) -> bytes:
    objs: list[bytes] = []
    n_pages = len(pages)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n_pages))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, lines in enumerate(pages):
        content_obj = 5 + 2 * i
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_obj} 0 R >>".encode()
        )
        stream = "BT /F1 12 Tf 72 720 Td 16 TL\n" + "\n".join(f"({l}) Tj T*" for l in lines) + "\nET"
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream".encode())
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def test_pdf_text_and_numbered_headings(tmp_path):
    f = tmp_path / "a.pdf"
    f.write_bytes(
        _make_pdf(
            [
                ["1. Introduction", "", "MQTT is a lightweight protocol.", "It uses a broker."],
                ["2. QoS", "", "QoS 0 means at most once."],
            ]
        )
    )
    m = parse(f)
    assert m.file_type == "pdf" and m.page_count == 2
    assert [h.text for h in m.headings] == ["1. Introduction", "2. QoS"]
    assert m.sections[1].page == 2
    assert "lightweight protocol" in m.sections[0].text


def test_pdf_without_headings_splits_by_page(tmp_path):
    f = tmp_path / "b.pdf"
    f.write_bytes(_make_pdf([["just some text here."], ["more text on page two."]]))
    m = parse(f)
    assert [s.title for s in m.sections] == ["페이지 1", "페이지 2"]
    assert m.warnings


def test_pdf_with_no_text_layer_is_empty_document(tmp_path):
    f = tmp_path / "scan.pdf"
    f.write_bytes(_make_pdf([[]]))
    with pytest.raises(EmptyDocument):
        parse(f)


# ------------------------------------------------------------ error cases
def test_corrupted_files_raise_readable_parsing_error(tmp_path):
    for name in ("bad.pdf", "bad.docx", "bad.pptx"):
        f = tmp_path / name
        f.write_bytes(b"this is not a real document")
        with pytest.raises(DocumentParsingError) as e:
            parse(f)
        assert "Traceback" not in e.value.message


def test_legacy_ppt_gives_actionable_message(tmp_path):
    f = tmp_path / "old.ppt"
    f.write_bytes(b"\xd0\xcf\x11\xe0legacy")
    with pytest.raises(DocumentParsingError) as e:
        parse(f)
    assert "PPTX" in e.value.message
