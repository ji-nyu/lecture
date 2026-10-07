"""STAGE 8: ClaudeProvider (Anthropic Messages API + pptx Agent Skill + Files API).

Nothing here talks to Anthropic. The real request/response code runs against a local fake server that
returns the shapes the official docs describe (`bash_code_execution_tool_result` ->
`bash_code_execution_result` -> `content[].file_id`, `/v1/files/{id}` and `/v1/files/{id}/content`,
`stop_reason: pause_turn` + `container.id`). One test starts the real detached runner process.
"""

import io
import json
import re
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import pytest
from pptx import Presentation

import app
from app.errors import PresentationGenerationFailed, PresentationProviderUnavailable
from app.models.presentation import JobState, PresentationRequest, SlideBrief
from app.providers import claude as claude_module
from app.providers.claude import (
    CODE_EXECUTION_TOOL,
    SKILL,
    ClaudeProvider,
    build_instructions,
    build_prompt,
    run_job,
)

KEY = "test-key-not-a-real-credential"
APP_DIR = Path(app.__file__).parent
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def request(n=4):
    slides = [SlideBrief(slide_number=i, title=f"슬라이드 제목 {i}", slide_type="concept") for i in range(1, n + 1)]
    return PresentationRequest(
        project_id="a" * 32, title="MQTT 입문", final_prompt="FINAL PROMPT BODY", prompt_hash="f" * 64,
        slide_count=n, slides=slides, duration_minutes=60,
    )


def deck(n=4) -> bytes:
    prs = Presentation()
    for i in range(n):
        prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = f"s{i}"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ a fake Anthropic API
def message(*blocks, stop="end_turn", container="cont_1"):
    return 200, {
        "id": "msg_1", "type": "message", "role": "assistant", "model": "m", "content": list(blocks),
        "stop_reason": stop, "stop_sequence": None,
        "usage": {"input_tokens": 11, "output_tokens": 22},
        "container": {"id": container, "expires_at": "2030-01-01T00:00:00Z"},
    }


def text(t="done"):
    return {"type": "text", "text": t}


def files_block(*file_ids):
    return {
        "type": "bash_code_execution_tool_result", "tool_use_id": "srvtoolu_1",
        "content": {
            "type": "bash_code_execution_result", "stdout": "", "stderr": "", "return_code": 0,
            "content": [{"type": "bash_code_execution_output", "file_id": f} for f in file_ids],
        },
    }


def api_error(status, kind, retry_ms=None):
    return status, {"type": "error", "error": {"type": kind, "message": "SECRET-DETAIL-FROM-API"}}, retry_ms


class FakeApi:
    def __init__(self):
        self.script = []  # answers to POST /v1/messages, in order; the last one repeats
        self.files = {}  # file_id -> (filename, bytes)
        self.requests = []  # (method, path, headers, json body)
        self._server = None

    def start(self):
        api = self

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def _send(self, status, payload, extra=None, raw=False):
                body = payload if raw else json.dumps(payload).encode()
                self.send_response(status)
                self.send_header("content-type", "application/octet-stream" if raw else "application/json")
                self.send_header("content-length", str(len(body)))
                for k, v in (extra or {}).items():
                    self.send_header(k, v)
                self.end_headers()
                self.wfile.write(body)

            def do_POST(self):
                length = int(self.headers.get("content-length") or 0)
                body = json.loads(self.rfile.read(length) or b"{}")
                api.requests.append(("POST", self.path, {k.lower(): v for k, v in self.headers.items()}, body))
                answer = api.script.pop(0) if len(api.script) > 1 else api.script[0]
                status, payload = answer[0], answer[1]
                retry = answer[2] if len(answer) > 2 and answer[2] else None
                self._send(status, payload, {"retry-after-ms": str(retry)} if retry else None)

            def do_GET(self):
                api.requests.append(("GET", self.path, {k.lower(): v for k, v in self.headers.items()}, None))
                m = re.match(r"^/v1/files/([^/?]+)(/content)?", self.path)
                if not m or m.group(1) not in api.files:
                    return self._send(404, {"type": "error", "error": {"type": "not_found_error", "message": "x"}})
                name, data = api.files[m.group(1)]
                if m.group(2):
                    return self._send(200, data, raw=True)
                self._send(200, {
                    "id": m.group(1), "type": "file", "filename": name, "mime_type": PPTX,
                    "size_bytes": len(data), "created_at": "2026-01-01T00:00:00Z", "downloadable": True,
                })

        self._server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        threading.Thread(target=self._server.serve_forever, daemon=True).start()
        return self

    @property
    def url(self):
        return f"http://127.0.0.1:{self._server.server_address[1]}"

    def stop(self):
        self._server.shutdown()
        self._server.server_close()

    def posts(self):
        return [r for r in self.requests if r[0] == "POST"]


@pytest.fixture
def api(monkeypatch):
    for var in ("HTTP_PROXY", "HTTPS_PROXY", "http_proxy", "https_proxy", "ALL_PROXY", "all_proxy"):
        monkeypatch.delenv(var, raising=False)
    monkeypatch.setenv("NO_PROXY", "127.0.0.1,localhost")
    fake = FakeApi().start()
    fake.script = [message(text(), files_block("file_1"))]
    fake.files = {"file_1": ("deck.pptx", deck(4))}
    yield fake
    fake.stop()


def in_thread(job_dir: Path) -> None:
    """Launcher for tests: the real runner code, in a thread (the SDK builds its client from the env)."""
    threading.Thread(target=run_job, args=(job_dir,), daemon=True).start()


@pytest.fixture
def provider(tmp_path, api, monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", KEY)
    monkeypatch.setenv("ANTHROPIC_BASE_URL", api.url)
    return ClaudeProvider(tmp_path / "jobs", api_key=KEY, model="claude-test-model", max_tokens=1234,
                          base_url=api.url, timeout_seconds=60, launcher=in_thread)


def finish(p, job_id, seconds=30):
    end = time.time() + seconds
    while time.time() < end:
        st = p.get_status(job_id)
        if st.state in (JobState.completed, JobState.failed):
            return st
        time.sleep(0.05)
    raise AssertionError("the job did not finish")


def run(p, n=4):
    job = p.create_presentation(request(n)).job_id
    return job, finish(p, job)


# ------------------------------------------------------------------ availability
def test_without_a_key_it_is_not_configured_and_names_the_variable(tmp_path):
    p = ClaudeProvider(tmp_path, api_key=None)
    a = p.availability()
    assert a.configured is False and "ANTHROPIC_API_KEY" in a.message
    with pytest.raises(PresentationProviderUnavailable) as e:
        p.create_presentation(request())
    assert "ANTHROPIC_API_KEY" in str(e.value)
    assert not list(tmp_path.iterdir())  # nothing was created


def test_an_empty_model_and_a_bad_base_url_are_not_configured(tmp_path):
    assert "CLAUDE_MODEL" in ClaudeProvider(tmp_path, api_key=KEY, model=" ").availability().message
    assert "CLAUDE_BASE_URL" in ClaudeProvider(tmp_path, api_key=KEY, base_url="ftp://x").availability().message


def test_a_missing_sdk_is_reported(tmp_path, monkeypatch):
    monkeypatch.setattr(claude_module.importlib.util, "find_spec", lambda name: None)
    a = ClaudeProvider(tmp_path, api_key=KEY).availability()
    assert a.configured is False and "anthropic" in a.message


def test_a_configured_provider_is_available(tmp_path):
    p = ClaudeProvider(tmp_path, api_key=KEY)
    assert p.availability().configured is True and p.name == "claude" and p.is_mock is False


def test_the_key_never_shows_in_repr(tmp_path):
    p = ClaudeProvider(tmp_path, api_key=KEY)
    assert KEY not in repr(p) and KEY not in str(p.__dict__.get("model"))


# ------------------------------------------------------------------ the request (documented shape)
def test_the_request_uses_the_pptx_skill_and_code_execution(provider, api):
    job, st = run(provider)
    assert st.state == JobState.completed
    (_, path, headers, body), = api.posts()
    assert path.startswith("/v1/messages")
    assert body["model"] == "claude-test-model" and body["max_tokens"] == 1234
    assert body["container"] == {"skills": [{"type": "anthropic", "skill_id": "pptx", "version": "latest"}]}
    assert body["tools"] == [{"type": "code_execution_20250825", "name": "code_execution"}]
    assert SKILL == body["container"]["skills"][0] and CODE_EXECUTION_TOOL == body["tools"][0]
    assert headers["x-api-key"] == KEY and headers["anthropic-version"] == "2023-06-01"
    assert "anthropic-beta" not in headers  # the current docs' examples use none


def test_the_prompt_is_sent_verbatim_and_the_system_text_fixes_the_slide_list(provider, api):
    run(provider, 3)
    body = api.posts()[0][3]
    user = body["messages"][0]
    assert user["role"] == "user" and "FINAL PROMPT BODY" in json.dumps(user["content"], ensure_ascii=False)
    assert "정확히 3장" in json.dumps(user["content"], ensure_ascii=False)
    system = body["system"] if isinstance(body["system"], str) else json.dumps(body["system"], ensure_ascii=False)
    assert "정확히 3장" in system and "1. 슬라이드 제목 1" in system and "3. 슬라이드 제목 3" in system


def test_the_instructions_keep_the_renderer_in_its_role():
    text_ = build_instructions(request(2))
    assert "구조" in text_ and "맡지 않는" in text_ and "인터넷이 없습니다" in text_
    assert "FINAL PROMPT BODY" not in text_  # the lecture is the user message, not the system text
    assert build_prompt(request(2)).startswith("FINAL PROMPT BODY")


# ------------------------------------------------------------------ success
def test_a_generated_file_is_downloaded_and_counted(provider, api):
    job, st = run(provider, 4)
    assert st.state == JobState.completed and st.progress == 100
    res = provider.get_result(job)
    assert res.slide_count == 4 and res.file_name.endswith(".pptx") and res.content_type == PPTX
    data = provider.download(job).data
    assert len(Presentation(io.BytesIO(data)).slides) == 4 and res.size_bytes == len(data)
    assert any(r[1].startswith("/v1/files/file_1/content") for r in api.requests)


def test_a_status_poll_before_the_runner_finishes_reports_running(provider, api):
    api.script = [message(text(), files_block("file_1"))]
    gate = threading.Event()
    real = claude_module._generate

    def slow(*a, **k):
        gate.wait(10)
        return real(*a, **k)

    claude_module._generate = slow
    try:
        job = provider.create_presentation(request()).job_id
        time.sleep(0.3)
        assert provider.get_status(job).state in (JobState.queued, JobState.running)
        with pytest.raises(PresentationGenerationFailed):
            provider.download(job)  # not finished
        gate.set()
        assert finish(provider, job).state == JobState.completed
    finally:
        claude_module._generate = real
        gate.set()


def test_pause_turn_is_continued_in_the_same_container(provider, api):
    api.script = [
        message(text("working"), stop="pause_turn", container="cont_42"),
        message(text("done"), files_block("file_1"), container="cont_42"),
    ]
    job, st = run(provider)
    assert st.state == JobState.completed
    first, second = (r[3] for r in api.posts())
    assert second["container"]["id"] == "cont_42" and second["container"]["skills"] == [SKILL]
    assert [m["role"] for m in second["messages"]] == ["user", "assistant"]
    assert second["messages"][1]["content"][0]["text"] == "working"
    assert "id" not in first["container"]


def test_the_newest_pptx_wins_and_other_files_are_skipped(provider, api):
    api.files = {"file_a": ("old.pptx", deck(2)), "file_b": ("notes.txt", b"hi"), "file_c": ("new.pptx", deck(5))}
    api.script = [message(text(), files_block("file_a", "file_b", "file_c"))]
    job, st = run(provider, 5)
    assert provider.get_result(job).slide_count == 5
    api.script = [message(text(), files_block("file_a", "file_b"))]
    api.files.pop("file_c")
    job2, st2 = run(provider, 2)
    assert provider.get_result(job2).slide_count == 2


def test_a_slide_count_that_differs_is_reported_not_hidden(provider, api):
    api.files = {"file_1": ("deck.pptx", deck(3))}
    job, st = run(provider, 6)  # the request says 6, Claude delivered 3
    assert st.state == JobState.completed and provider.get_result(job).slide_count == 3


# ------------------------------------------------------------------ failures
@pytest.mark.parametrize("answer, expected", [
    (api_error(401, "authentication_error"), "API 키를 거부"),
    (api_error(403, "permission_error"), "쓸 수 없습니다"),
    (api_error(404, "not_found_error"), "CLAUDE_MODEL"),
    (api_error(400, "invalid_request_error"), "크레딧"),
    (api_error(429, "rate_limit_error", retry_ms=1), "요청 한도"),
    (api_error(500, "api_error", retry_ms=1), "일시적으로"),
])
def test_api_errors_become_short_messages_without_api_text(provider, api, answer, expected):
    api.script = [answer]
    job, st = run(provider)
    assert st.state == JobState.failed and expected in st.message
    for forbidden in ("SECRET-DETAIL-FROM-API", KEY, "Traceback", "anthropic."):
        assert forbidden not in st.message
    with pytest.raises(PresentationGenerationFailed) as e:
        provider.get_result(job)
    assert "SECRET-DETAIL-FROM-API" not in str(e.value)


def test_no_file_and_the_reason_for_it(provider, api):
    api.script = [message(text("I could not"))]
    assert run(provider)[1].message.startswith("Claude가 PPTX 파일을 만들지 못했습니다")
    api.script = [message(text("cut"), stop="max_tokens")]
    assert "CLAUDE_MAX_TOKENS" in run(provider)[1].message
    api.script = [message(text("no"), stop="refusal")]
    assert "처리하지 않았습니다" in run(provider)[1].message


def test_a_file_that_is_not_a_pptx_fails(provider, api):
    api.files = {"file_1": ("notes.txt", b"hello")}
    st = run(provider)[1]
    assert st.state == JobState.failed and "PPTX가 없습니다" in st.message


def test_a_broken_pptx_fails(provider, api):
    api.files = {"file_1": ("deck.pptx", b"this is not a zip file")}
    st = run(provider)[1]
    assert st.state == JobState.failed and "올바른 PPTX가 아닙니다" in st.message


def test_endless_pause_turns_stop(provider, api):
    api.script = [message(text("again"), stop="pause_turn")]
    st = run(provider)[1]
    assert st.state == JobState.failed and "정해진 횟수" in st.message
    assert len(api.posts()) == claude_module.MAX_TURNS


def test_a_job_past_its_deadline_is_stopped(tmp_path, api, monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", KEY)
    monkeypatch.setenv("ANTHROPIC_BASE_URL", api.url)
    p = ClaudeProvider(tmp_path / "jobs", api_key=KEY, base_url=api.url, timeout_seconds=0, launcher=in_thread)
    st = run(p)[1]
    assert st.state == JobState.failed and "제한 시간" in st.message and not api.posts()


def test_a_dead_runner_is_detected_by_its_missing_heartbeat(tmp_path):
    p = ClaudeProvider(tmp_path / "jobs", api_key=KEY, launcher=lambda d: None)  # nothing ever runs
    job = p.create_presentation(request()).job_id
    assert p.get_status(job).state == JobState.queued
    job_json = tmp_path / "jobs" / job / "job.json"
    data = json.loads(job_json.read_text(encoding="utf-8"))
    data["created_at"] = time.time() - 3600
    job_json.write_text(json.dumps(data), encoding="utf-8")
    st = p.get_status(job)
    assert st.state == JobState.failed and "중단" in st.message


def test_a_launcher_that_cannot_start_is_unavailable(tmp_path):
    def broken(_):
        raise OSError("SECRET-PATH-DETAIL")

    p = ClaudeProvider(tmp_path, api_key=KEY, launcher=broken)
    with pytest.raises(PresentationProviderUnavailable) as e:
        p.create_presentation(request())
    assert "SECRET-PATH-DETAIL" not in str(e.value)


def test_job_ids_cannot_escape_the_jobs_directory(provider):
    for bad in ("../x", "cld-../../etc", "", "gsk-" + "a" * 32, "cld-" + "g" * 32):
        with pytest.raises(PresentationGenerationFailed):
            provider.get_status(bad)
    with pytest.raises(PresentationGenerationFailed):
        provider.get_status("cld-" + "a" * 32)  # well formed, but unknown


# ------------------------------------------------------------------ secrets
def test_the_key_is_never_written_to_the_job_directory(provider, api):
    job, _ = run(provider)
    for f in (provider.root / job).rglob("*"):
        if f.is_file():
            assert KEY.encode() not in f.read_bytes(), f.name


def test_the_key_is_redacted_in_the_error_file(tmp_path, monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", KEY)

    class Boom:
        class messages:
            @staticmethod
            def create(**kw):
                raise RuntimeError(f"leaked {KEY} here")

    d = tmp_path / "job"
    d.mkdir()
    (d / "job.json").write_text(json.dumps({"timeout_seconds": 60}), encoding="utf-8")
    (d / "args.json").write_text(json.dumps({"model": "m", "max_tokens": 10, "system": "s", "prompt": "p"}), encoding="utf-8")
    run_job(d, client=Boom())
    assert KEY not in (d / "error.txt").read_text(encoding="utf-8")
    assert json.loads((d / "result.json").read_text(encoding="utf-8"))["error_kind"] == "unexpected"


def test_the_source_holds_no_key_and_no_endpoint():
    src = (APP_DIR / "providers" / "claude.py").read_text(encoding="utf-8")
    assert "sk-ant" not in src and "api.anthropic.com" not in src and KEY not in src


def test_the_child_environment_carries_only_the_anthropic_key(tmp_path, monkeypatch):
    monkeypatch.setenv("LLM_API_KEY", "other-secret")
    monkeypatch.setenv("OPENAI_API_KEY", "other-secret")
    monkeypatch.setenv("ANTHROPIC_AUTH_TOKEN", "other-secret")
    monkeypatch.setenv("ANTHROPIC_BASE_URL", "http://elsewhere.invalid")
    env = ClaudeProvider(tmp_path, api_key=KEY)._child_env()
    assert env["ANTHROPIC_API_KEY"] == KEY and "ANTHROPIC_BASE_URL" not in env
    assert not {"LLM_API_KEY", "OPENAI_API_KEY", "ANTHROPIC_AUTH_TOKEN"} & set(env)
    assert ClaudeProvider(tmp_path, api_key=KEY, base_url="http://127.0.0.1:1")._child_env()["ANTHROPIC_BASE_URL"] == "http://127.0.0.1:1"


def test_the_failure_is_logged_once_and_without_the_key(provider, api, caplog):
    api.script = [api_error(401, "authentication_error")]
    with caplog.at_level("WARNING", logger="ailecturegen"):
        job, _ = run(provider)
        provider.get_status(job)
        provider.get_status(job)
    lines = [r.getMessage() for r in caplog.records if "failed" in r.getMessage()]
    assert len(lines) == 1 and "auth" in lines[0] and KEY not in caplog.text


# ------------------------------------------------------------------ the real detached runner
def test_the_detached_runner_process_works_end_to_end(tmp_path, api, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_BASE_URL", raising=False)
    api.files = {"file_1": ("deck.pptx", deck(3))}
    p = ClaudeProvider(tmp_path / "jobs", api_key=KEY, model="claude-test-model", base_url=api.url, timeout_seconds=120)
    job, st = run(p, 3)  # default launcher: `python -m app.providers.claude <job_dir>` in a new process
    assert st.state == JobState.completed, st.message
    assert p.get_result(job).slide_count == 3
    (_, _, headers, body), = api.posts()
    assert headers["x-api-key"] == KEY and body["model"] == "claude-test-model"
    for f in (p.root / job).rglob("*"):
        if f.is_file():
            assert KEY.encode() not in f.read_bytes(), f.name
