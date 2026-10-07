"""Lecture video: PPT slides + spoken scripts -> one downloadable MP4."""

from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

from app.config import Settings
from app.llm.mock import MockLLMClient
from app.main import create_app

from app.services.lecture_video import (
    HOLD_AFTER_SPEECH,
    VideoPipeline,
    audio_seconds,
    ffmpeg_exe,
    mux_clips,
    run_job,
    write_silence,
)

from test_stage7_api import generated


class FakePipeline(VideoPipeline):
    def render_slides(self, pptx: Path, dest: Path):
        dest.mkdir(parents=True, exist_ok=True)
        deck = Presentation(str(pptx))
        paths = []
        for i, _ in enumerate(deck.slides, 1):
            path = dest / f"slide-{i:03d}.png"
            Image.new("RGB", (64, 64), (255, 255, 255)).save(path)
            paths.append(path)
        return paths, None

    def speak(self, text, dest: Path, min_seconds: float = 0):
        return write_silence(dest, 0.25)

    def mux(self, pairs, dest: Path):
        dest.write_bytes(b"FAKEMP4" + bytes([len(pairs)]) + b"\x00" * 40)
        return dest, 0.25 * len(pairs)

    def prepend(self, intro: Path, lecture: Path) -> float:
        lecture.write_bytes(lecture.read_bytes() + b"+INTRO")
        return 1.5


def video_client(tmp_path: Path) -> TestClient:
    pipe = FakePipeline()
    settings = Settings(data_dir=tmp_path, max_upload_mb=5, cors_origins=(), enrichment_provider="mock")
    return TestClient(
        create_app(
            settings,
            content_client=MockLLMClient(),
            video_pipeline=pipe,
            video_launcher=lambda d: run_job(d, pipe),
        ),
        raise_server_exceptions=False,
    )


def test_video_needs_a_finished_ppt(tmp_path: Path):
    c = video_client(tmp_path)
    from test_stage4_api import planned

    pid = planned(c)
    assert c.post(f"/projects/{pid}/slides").status_code == 200
    r = c.post(f"/projects/{pid}/video")
    assert r.status_code == 409
    assert r.json()["error"]["code"] in {"PresentationNotReady", "VideoNotReady"}


def test_video_is_built_from_ppt_and_scripts(tmp_path: Path):
    c = video_client(tmp_path)
    pid, _ = generated(c)
    assert c.get(f"/projects/{pid}/scripts").status_code == 200
    r = c.post(f"/projects/{pid}/video")
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["state"] == "completed" and body["has_result"] is True
    assert body["slide_count"] >= 1
    assert c.get(f"/projects/{pid}").json()["has_video"] is True
    dl = c.get(f"/projects/{pid}/video/download")
    assert dl.status_code == 200 and dl.content.startswith(b"FAKEMP4")
    again = c.post(f"/projects/{pid}/video")
    assert again.status_code == 200 and again.json()["state"] == "completed"


def test_render_paints_the_pptx_picture_not_a_text_poster(tmp_path: Path):
    from io import BytesIO

    from pptx.util import Inches

    from app.services.lecture_video import _rasterize_pptx

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    buf = BytesIO()
    Image.new("RGB", (80, 45), (220, 30, 30)).save(buf, "PNG")
    buf.seek(0)
    slide.shapes.add_picture(buf, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    pptx = tmp_path / "deck.pptx"
    prs.save(pptx)
    images = _rasterize_pptx(pptx, tmp_path / "frames")
    assert len(images) == 1
    pix = Image.open(images[0]).getpixel((20, 20))
    assert pix[0] > 180 and pix[1] < 80


def test_rasterize_keeps_large_text_and_ovals(tmp_path: Path):
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    from pptx.util import Inches, Pt

    from app.services.lecture_video import HEIGHT, WIDTH, _rasterize_pptx

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(2.5))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "AIoT"
    run.font.size = Pt(180)
    run.font.color.rgb = RGBColor(181, 83, 58)
    oval = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(10), Inches(1), Inches(2), Inches(2))
    oval.fill.solid()
    oval.fill.fore_color.rgb = RGBColor(0, 140, 40)
    pptx = tmp_path / "deck.pptx"
    prs.save(pptx)
    images = _rasterize_pptx(pptx, tmp_path / "frames")
    img = Image.open(images[0])
    sx, sy = WIDTH / 13.333, HEIGHT / 7.5
    cx, cy = int(11 * sx), int(2 * sy)
    corner = img.getpixel((int(10 * sx) + 2, int(1 * sy) + 2))
    center = img.getpixel((cx, cy))
    assert center[1] > 80 and center[0] < 80
    assert corner[1] < 80 or corner[0] > 80
    region = img.crop((int(0.5 * sx), int(0.5 * sy), int(8.5 * sx), int(3.0 * sy)))
    terracotta = sum(1 for px in region.getdata() if px[0] > 120 and px[1] < 120 and px[2] < 120)
    assert terracotta > 400


def test_rasterize_keeps_per_run_colors(tmp_path: Path):
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    from app.services.lecture_video import HEIGHT, WIDTH, _rasterize_pptx

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(10), Inches(2.5))
    p = box.text_frame.paragraphs[0]
    a = p.add_run()
    a.text, a.font.size, a.font.bold = "A", Pt(180), True
    a.font.color.rgb = RGBColor(181, 83, 58)
    i = p.add_run()
    i.text, i.font.size, i.font.bold = "I", Pt(180), True
    i.font.color.rgb = RGBColor(26, 39, 48)
    pptx = tmp_path / "deck.pptx"
    prs.save(pptx)
    img = Image.open(_rasterize_pptx(pptx, tmp_path / "frames")[0])
    sx, sy = WIDTH / 13.333, HEIGHT / 7.5
    region = img.crop((int(1 * sx), int(2 * sy), int(11 * sx), int(4.5 * sy)))
    terra = sum(1 for px in region.getdata() if px[0] > 140 and px[1] < 110)
    navy = sum(1 for px in region.getdata() if px[0] < 50 and px[1] < 60 and px[2] < 80)
    assert terra > 200 and navy > 200


def test_a_slide_stays_until_its_spoken_audio_ends(tmp_path: Path):
    try:
        exe = ffmpeg_exe()
    except RuntimeError:
        pytest.skip("ffmpeg가 없습니다")
    img = tmp_path / "slide.png"
    Image.new("RGB", (128, 72), (20, 40, 80)).save(img)
    short = write_silence(tmp_path / "short.wav", 0.7)
    long = write_silence(tmp_path / "long.wav", 1.5)
    out = tmp_path / "lecture.mp4"
    _, total = mux_clips([(img, short), (img, long)], out)
    assert out.is_file() and out.stat().st_size > 32
    expected = 0.7 + 1.5 + 2 * HOLD_AFTER_SPEECH
    assert abs(total - expected) < 0.15
    assert audio_seconds(out, exe) >= 0.7 + 1.5 - 0.1


def test_a_slide_advances_when_speech_ends(tmp_path: Path):
    try:
        exe = ffmpeg_exe()
    except RuntimeError:
        pytest.skip("ffmpeg가 없습니다")
    short = write_silence(tmp_path / "short.wav", 0.4)
    img = tmp_path / "slide.png"
    Image.new("RGB", (128, 72), (20, 40, 80)).save(img)
    out = tmp_path / "spoken.mp4"
    _, total = mux_clips([(img, short)], out)
    assert abs(total - (0.4 + HOLD_AFTER_SPEECH)) < 0.15
    assert audio_seconds(out, exe) < 0.9


def test_new_ppt_clears_the_video(tmp_path: Path):
    c = video_client(tmp_path)
    pid, _ = generated(c)
    assert c.get(f"/projects/{pid}/scripts").status_code == 200
    assert c.post(f"/projects/{pid}/video").status_code == 200
    assert c.get(f"/projects/{pid}").json()["has_video"] is True
    assert c.post(f"/projects/{pid}/presentation").status_code == 200
    assert c.get(f"/projects/{pid}").json()["has_video"] is False
    assert c.get(f"/projects/{pid}/video/download").status_code == 409
