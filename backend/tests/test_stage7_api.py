"""STAGE 7 over HTTP: prompt -> provider -> job -> file, states, failures, invalidation, restart."""

import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app.config import Settings
from app.errors import PresentationProviderUnavailable
from app.llm.mock import MockLLMClient
from app.main import create_app
from app.models.presentation import DownloadedFile, JobState, ProviderResult, ProviderStatus
from app.models.project import PresentationStatus
from app.providers.base import PresentationProvider, ProviderAvailability
from app.providers.mock import MockPresentationProvider
from app.services.presentation_service import _download_name

from plan_helpers import BASE, CASE_B
from test_stage6a_api import approved, new_project, project, put_profile

PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def make(tmp_path, provider=None, **settings):
    s = Settings(data_dir=tmp_path, max_upload_mb=5, cors_origins=(), enrichment_provider="mock", **settings)
    return TestClient(
        create_app(s, content_client=MockLLMClient(), presentation_provider=provider),
        raise_server_exceptions=False,
    )


@pytest.fixture()
def client(tmp_path):
    return make(tmp_path)


def ready(client, **opts):
    """Approved lecture with a built prompt (status ready_for_presentation)."""
    pid = approved(client, **{**CASE_B, **opts})
    r = client.post(f"/projects/{pid}/prompt")
    assert r.status_code == 200, r.text
    assert project(client, pid)["presentation_status"] == "ready_for_presentation"
    return pid


def generated(client, **opts):
    pid = ready(client, **opts)
    r = client.post(f"/projects/{pid}/presentation")
    assert r.status_code == 200, r.text
    return pid, r.json()


def stored(tmp_path, pid):
    return json.loads((tmp_path / "projects" / pid / "project.json").read_text(encoding="utf-8"))


def deck_dir(tmp_path, pid):
    return tmp_path / "projects" / pid / "presentation"


class Fake(PresentationProvider):
    """A configurable provider for the cases the mock does not cover."""

    name = "fake"
    is_mock = False

    def __init__(self, *, file_name="deck.pptx", data=b"PK\x03\x04 fake deck", slide_count=None,
                 create_exc=None, status_exc=None, download_exc=None, on_create=None, states=None):
        self.file_name, self.data, self.slide_count = file_name, data, slide_count
        self.create_exc, self.status_exc, self.download_exc = create_exc, status_exc, download_exc
        self.on_create = on_create
        self.states = list(states or [JobState.completed])
        self.calls = 0

    def availability(self):
        return ProviderAvailability(configured=True)

    def create_presentation(self, request):
        if self.on_create:
            self.on_create()
        if self.create_exc:
            raise self.create_exc
        self.calls += 1
        return ProviderStatus(job_id=f"fake-{self.calls}", state=self.states[0], progress=0)

    def get_status(self, job_id):
        if self.status_exc:
            raise self.status_exc
        state = self.states.pop(0) if len(self.states) > 1 else self.states[0]
        return ProviderStatus(job_id=job_id, state=state, progress=100 if state == JobState.completed else 10)

    def get_result(self, job_id):
        return ProviderResult(job_id=job_id, file_name=self.file_name, content_type=PPTX,
                              size_bytes=len(self.data), slide_count=self.slide_count)

    def download(self, job_id):
        if self.download_exc:
            raise self.download_exc
        return DownloadedFile(file_name=self.file_name, content_type=PPTX, data=self.data)


# ------------------------------------------------------------------ happy path
def test_the_prompt_becomes_a_presentation(client, tmp_path):
    pid = ready(client)
    slides = client.get(f"/projects/{pid}/slides").json()
    plan = client.get(f"/projects/{pid}/plan").json()
    prompt = client.get(f"/projects/{pid}/prompt").json()

    r = client.post(f"/projects/{pid}/presentation")
    assert r.status_code == 200, r.text
    s = r.json()
    assert s["state"] == "completed" and s["project_status"] == "completed" and s["has_result"] is True
    assert s["provider"] == "mock" and s["is_mock"] is True and s["provider_configured"] is True
    assert s["progress"] == 100 and s["prompt_hash"] == prompt["prompt_hash"] and s["job_id"].startswith("mock-")

    pr = project(client, pid)
    assert pr["presentation_status"] == "completed" and pr["has_presentation"] is True
    assert pr["presentation_provider"] == "mock" and pr["error_message"] is None

    res = client.get(f"/projects/{pid}/presentation/result").json()
    assert res["is_mock"] is True and res["prompt_hash"] == prompt["prompt_hash"]
    assert res["slide_count"] == slides["slide_count"] and res["size_bytes"] > 0 and len(res["sha256"]) == 64
    assert "MOCK" in res["file_name"] and res["file_name"].endswith(".pptx") and "디자인" in res["note"]

    # the lecture itself was not touched
    assert client.get(f"/projects/{pid}/slides").json() == slides
    assert client.get(f"/projects/{pid}/plan").json() == plan
    assert client.get(f"/projects/{pid}/prompt").json() == prompt


def test_the_downloaded_file_has_the_planned_slides_in_order(client):
    pid, _ = generated(client)
    slides = client.get(f"/projects/{pid}/slides").json()["slides"]
    r = client.get(f"/projects/{pid}/presentation/download")
    assert r.status_code == 200 and r.headers["content-type"] == PPTX
    assert "attachment" in r.headers["content-disposition"] and ".pptx" in r.headers["content-disposition"]
    prs = Presentation(io.BytesIO(r.content))
    assert [s.shapes.title.text for s in prs.slides] == [s["title"] for s in slides]
    res = client.get(f"/projects/{pid}/presentation/result").json()
    import hashlib
    assert hashlib.sha256(r.content).hexdigest() == res["sha256"] and len(r.content) == res["size_bytes"]


def test_the_file_is_kept_with_the_project_and_no_server_path_is_exposed(client, tmp_path):
    pid, status = generated(client)
    assert (deck_dir(tmp_path, pid) / "presentation.pptx").is_file()
    st = stored(tmp_path, pid)
    assert st["presentation_result"]["provider"] == "mock" and st["presentation_job"]["state"] == "completed"
    for r in (client.get(f"/projects/{pid}"), client.get(f"/projects/{pid}/presentation/status"),
              client.get(f"/projects/{pid}/presentation/result")):
        assert str(tmp_path) not in r.text and "presentation.pptx" not in r.text


def test_status_before_anything_was_requested(client):
    pid = ready(client)
    s = client.get(f"/projects/{pid}/presentation/status").json()
    assert s["state"] == "not_started" and s["has_result"] is False and s["job_id"] is None
    assert s["provider"] == "mock" and s["provider_configured"] is True
    assert client.get(f"/projects/{pid}/presentation/result").status_code == 409
    assert client.get(f"/projects/{pid}/presentation/download").status_code == 409


def test_regenerating_replaces_the_previous_file(client, tmp_path):
    pid, first = generated(client)
    r = client.post(f"/projects/{pid}/presentation")
    assert r.status_code == 200 and r.json()["state"] == "completed"
    assert r.json()["job_id"] != first["job_id"]
    assert client.get(f"/projects/{pid}/presentation/result").json()["job_id"] == r.json()["job_id"]
    assert len(list(deck_dir(tmp_path, pid).iterdir())) == 1
    assert client.get(f"/projects/{pid}/presentation/download").status_code == 200


def test_a_project_without_enrichment_and_without_any_llm_still_gets_a_deck(tmp_path):
    s = Settings(data_dir=tmp_path, max_upload_mb=5, cors_origins=())  # ENRICHMENT_PROVIDER=none
    c = TestClient(create_app(s), raise_server_exceptions=False)
    pid = approved(c, **CASE_B)
    assert c.post(f"/projects/{pid}/prompt").status_code == 200
    assert c.post(f"/projects/{pid}/presentation").json()["state"] == "completed"


def test_everything_survives_a_server_restart(tmp_path):
    c = make(tmp_path)
    pid, _ = generated(c)
    body = c.get(f"/projects/{pid}/presentation/download").content
    again = make(tmp_path)
    s = again.get(f"/projects/{pid}/presentation/status").json()
    assert s["state"] == "completed" and s["has_result"] is True
    assert again.get(f"/projects/{pid}/presentation/download").content == body
    assert again.get(f"/projects/{pid}").json()["presentation_status"] == "completed"


# ------------------------------------------------------------------ slow jobs (polling)
def test_a_slow_job_is_collected_by_polling(tmp_path):
    c = make(tmp_path, MockPresentationProvider(tmp_path / "jobs", polls_until_done=2))
    pid = ready(c)
    s = c.post(f"/projects/{pid}/presentation").json()
    assert s["state"] == "queued" and s["project_status"] == "generating" and s["has_result"] is False
    assert c.get(f"/projects/{pid}/presentation/result").status_code == 409
    assert c.get(f"/projects/{pid}/presentation/download").status_code == 409

    s = c.get(f"/projects/{pid}/presentation/status").json()
    assert s["state"] == "running" and s["progress"] == 50 and s["project_status"] == "generating"
    s = c.get(f"/projects/{pid}/presentation/status").json()
    assert s["state"] == "completed" and s["project_status"] == "completed" and s["has_result"] is True
    assert c.get(f"/projects/{pid}/presentation/download").status_code == 200
    # polling a finished job does not ask the provider again
    assert c.get(f"/projects/{pid}/presentation/status").json()["state"] == "completed"


def test_starting_twice_while_generating_is_refused(tmp_path):
    c = make(tmp_path, MockPresentationProvider(tmp_path / "jobs", polls_until_done=3))
    pid = ready(c)
    assert c.post(f"/projects/{pid}/presentation").status_code == 200
    r = c.post(f"/projects/{pid}/presentation")
    assert r.status_code == 409 and r.json()["error"]["code"] == "PresentationInProgress"


def test_an_interrupted_start_can_be_started_again(tmp_path):
    c = make(tmp_path)
    pid = ready(c)
    p = c.app.state.store.get(pid)  # the server died right after marking the project as generating
    p.presentation_status = PresentationStatus.generating
    c.app.state.store.save(p)
    assert c.get(f"/projects/{pid}/presentation/status").json()["state"] == "not_started"
    assert c.post(f"/projects/{pid}/presentation").json()["state"] == "completed"


def test_a_job_started_by_another_provider_cannot_be_continued(tmp_path):
    c = make(tmp_path, MockPresentationProvider(tmp_path / "jobs", polls_until_done=5))
    pid = ready(c)
    assert c.post(f"/projects/{pid}/presentation").json()["state"] == "queued"
    c.app.state.presentation_provider = Fake()  # settings changed, server restarted (another provider name)
    s = c.get(f"/projects/{pid}/presentation/status").json()
    assert s["state"] == "failed" and "다른 공급자" in s["message"] and s["project_status"] == "failed"


# ------------------------------------------------------------------ preconditions
def test_an_approved_lecture_builds_the_prompt_then_the_presentation(client):
    pid = approved(client, **CASE_B)  # approved, but the prompt was not built yet
    r = client.post(f"/projects/{pid}/presentation")
    assert r.status_code == 200, r.text
    assert r.json()["state"] == "completed" and project(client, pid)["has_prompt"] is True
    assert project(client, pid)["has_presentation"] is True

    pid2 = new_project(client)
    assert client.post(f"/projects/{pid2}/presentation").status_code == 409


def test_unknown_project(client):
    for method, path in (("post", ""), ("get", "/status"), ("get", "/result"), ("get", "/download")):
        r = getattr(client, method)(f"/projects/{'0' * 32}/presentation{path}")
        assert r.status_code == 404 and r.json()["error"]["code"] == "ProjectNotFound"


def test_no_presentation_while_the_content_is_being_enriched(client):
    pid = ready(client)
    p = client.app.state.store.get(pid)
    p.presentation_status = PresentationStatus.enriching
    client.app.state.store.save(p)
    r = client.post(f"/projects/{pid}/presentation")
    assert r.status_code == 409 and r.json()["error"]["code"] == "PresentationNotAllowed"


# ------------------------------------------------------------------ failures
def test_a_failed_job_is_reported_and_can_be_retried(tmp_path):
    c = make(tmp_path, MockPresentationProvider(tmp_path / "jobs", fail_at="render"))
    pid = ready(c)
    s = c.post(f"/projects/{pid}/presentation").json()
    assert s["state"] == "failed" and s["project_status"] == "failed" and s["message"] and s["has_result"] is False
    pr = project(c, pid)
    assert pr["presentation_status"] == "failed" and pr["error_message"] and pr["has_presentation"] is False
    assert c.get(f"/projects/{pid}/presentation/download").status_code == 409
    assert c.get(f"/projects/{pid}/prompt").status_code == 200  # the lecture and prompt are intact

    c.app.state.presentation_provider = MockPresentationProvider(tmp_path / "jobs")  # the outage is over
    s = c.post(f"/projects/{pid}/presentation").json()
    assert s["state"] == "completed" and project(c, pid)["error_message"] is None


def test_a_provider_that_refuses_to_start_changes_nothing(tmp_path):
    c = make(tmp_path, MockPresentationProvider(tmp_path / "jobs", fail_at="create"))
    pid = ready(c)
    before = stored(tmp_path, pid)
    r = c.post(f"/projects/{pid}/presentation")
    assert r.status_code == 503 and r.json()["error"]["code"] == "PresentationProviderUnavailable"
    after = stored(tmp_path, pid)
    assert after["presentation_status"] == "ready_for_presentation"
    assert after["presentation_job"] is None and after["presentation_result"] is None
    assert after["final_prompt"] == before["final_prompt"]
    assert c.get(f"/projects/{pid}/presentation/status").json()["state"] == "not_started"


def test_a_failed_regeneration_keeps_the_earlier_presentation(tmp_path):
    c = make(tmp_path)
    pid, _ = generated(c)
    c.app.state.presentation_provider = MockPresentationProvider(tmp_path / "jobs", fail_at="create")
    assert c.post(f"/projects/{pid}/presentation").status_code == 503
    assert project(c, pid)["presentation_status"] == "completed"
    assert c.get(f"/projects/{pid}/presentation/download").status_code == 200


def test_genspark_mode_reports_not_configured_and_never_calls_the_mock(tmp_path):
    c = make(tmp_path, genspark_mode="genspark", genspark_cli_path=str(tmp_path / "no-such-gsk"))
    pid = ready(c)  # everything before the provider still works
    s = c.get(f"/projects/{pid}/presentation/status").json()
    assert s["state"] == "not_started" and s["provider"] == "genspark" and s["provider_configured"] is False
    assert "gsk" in s["provider_message"] and s["is_mock"] is False

    r = c.post(f"/projects/{pid}/presentation")
    assert r.status_code == 503 and r.json()["error"]["code"] == "PresentationProviderUnavailable"
    assert "gsk" in r.json()["error"]["message"]
    pr = project(c, pid)
    assert pr["presentation_status"] == "ready_for_presentation" and pr["has_presentation"] is False
    assert not (tmp_path / "mock_presentations").exists()  # no fallback to the mock behind the user's back


def test_the_default_setup_uses_the_mock_under_the_data_dir(client, tmp_path):
    pid, s = generated(client)
    assert s["provider"] == "mock"
    assert (tmp_path / "mock_presentations" / s["job_id"] / "job.json").is_file()


@pytest.mark.parametrize("kwargs, needle", [
    (dict(file_name="deck.exe"), "형식"),
    (dict(file_name="deck.pptx", data=b""), "빈 파일"),
])
def test_an_unusable_provider_file_fails_the_job(tmp_path, kwargs, needle):
    c = make(tmp_path, Fake(**kwargs))
    pid = ready(c)
    s = c.post(f"/projects/{pid}/presentation").json()
    assert s["state"] == "failed" and needle in s["message"] and s["has_result"] is False
    assert not deck_dir(tmp_path, pid).exists()


@pytest.mark.parametrize("where", ["create", "status", "download"])
def test_provider_crashes_do_not_leak_details(tmp_path, where):
    boom = RuntimeError("secret token sk-abc123 at C:\\srv\\keys\\x.txt")
    fake = Fake(create_exc=boom if where == "create" else None,
                status_exc=boom if where == "status" else None,
                download_exc=boom if where == "download" else None,
                states=[JobState.running, JobState.completed] if where == "status" else None)
    c = make(tmp_path, fake)
    pid = ready(c)
    r = c.post(f"/projects/{pid}/presentation")
    if where == "create":
        assert r.status_code == 502 and r.json()["error"]["code"] == "PresentationGenerationFailed"
        assert project(c, pid)["presentation_status"] == "ready_for_presentation"
    elif where == "status":
        r = c.get(f"/projects/{pid}/presentation/status")
        assert r.json()["state"] == "failed"
    else:
        assert r.json()["state"] == "failed"
    for resp in (r, c.get(f"/projects/{pid}"), c.get(f"/projects/{pid}/presentation/status")):
        assert "sk-abc123" not in resp.text and "C:\\srv" not in resp.text and "Traceback" not in resp.text


def test_a_temporary_provider_outage_while_polling_keeps_the_job(tmp_path):
    fake = Fake(states=[JobState.running], status_exc=PresentationProviderUnavailable("잠시 사용할 수 없습니다."))
    c = make(tmp_path, fake)
    pid = ready(c)
    assert c.post(f"/projects/{pid}/presentation").json()["state"] == "running"
    r = c.get(f"/projects/{pid}/presentation/status")
    assert r.status_code == 503 and "잠시" in r.json()["error"]["message"]
    assert project(c, pid)["presentation_status"] == "generating"
    fake.status_exc, fake.states = None, [JobState.completed]
    assert c.get(f"/projects/{pid}/presentation/status").json()["state"] == "completed"


def test_a_slide_count_mismatch_is_reported_not_hidden(tmp_path):
    c = make(tmp_path, Fake(slide_count=999))
    pid = ready(c)
    assert c.post(f"/projects/{pid}/presentation").json()["state"] == "completed"
    res = c.get(f"/projects/{pid}/presentation/result").json()
    assert "999" in res["note"] and "다릅니다" in res["note"] and res["is_mock"] is False


def test_the_lecture_changing_during_generation_discards_the_result(tmp_path):
    holder = {}

    def change_options():  # the professor edits the options while the provider is working
        r = holder["c"].put(f"/projects/{holder['pid']}/profile", json={**CASE_B, "duration_minutes": 40})
        assert r.status_code == 200

    fake = Fake(on_create=change_options)
    c = make(tmp_path, fake)
    holder["c"] = c
    holder["pid"] = pid = ready(c)
    r = c.post(f"/projects/{pid}/presentation")
    assert r.status_code == 409 and r.json()["error"]["code"] == "PresentationStale"
    pr = project(c, pid)
    assert pr["has_presentation"] is False and pr["has_prompt"] is False and pr["presentation_status"] == "analyzed"


# ------------------------------------------------------------------ invalidation
CHANGES = {
    "profile": (lambda c, pid: c.put(f"/projects/{pid}/profile", json={**BASE, **CASE_B, "duration_minutes": 40}), "analyzed"),
    "plan": (lambda c, pid: c.post(f"/projects/{pid}/plan"), "planned"),
    "slides": (lambda c, pid: c.post(f"/projects/{pid}/slides"), "planned"),
    "enrich": (lambda c, pid: c.post(f"/projects/{pid}/enrich"), None),  # enriched / enrichment_partial
    "analyze": (lambda c, pid: c.post(f"/projects/{pid}/analyze"), "analyzed"),
    "upload": (lambda c, pid: c.post(f"/projects/{pid}/upload",
                                     files={"file": ("n.md", io.BytesIO(b"# New\n\ntext text text\n"), "text/plain")}), "uploaded"),
}


@pytest.mark.parametrize("what", list(CHANGES))
def test_changing_the_lecture_withdraws_the_presentation(client, tmp_path, what):
    pid, _ = generated(client)
    assert deck_dir(tmp_path, pid).exists()
    change, expected = CHANGES[what]
    r = change(client, pid)
    assert r.status_code == 200, r.text
    pr = project(client, pid)
    assert pr["has_presentation"] is False and pr["has_prompt"] is False and pr["presentation_provider"] is None
    if expected:
        assert pr["presentation_status"] == expected
    else:
        assert pr["presentation_status"] in ("enriched", "enrichment_partial")
    assert client.get(f"/projects/{pid}/presentation/result").status_code == 409
    assert client.get(f"/projects/{pid}/presentation/download").status_code == 409
    assert client.get(f"/projects/{pid}/presentation/status").json()["state"] == "not_started"
    assert not deck_dir(tmp_path, pid).exists()  # the old file is gone too
    assert stored(tmp_path, pid)["presentation_job"] is None


def test_a_failed_presentation_is_withdrawn_like_a_finished_one(tmp_path):
    c = make(tmp_path, MockPresentationProvider(tmp_path / "jobs", fail_at="render"))
    pid = ready(c)
    c.post(f"/projects/{pid}/presentation")
    assert project(c, pid)["presentation_status"] == "failed"
    put_profile(c, pid, **{**CASE_B, "duration_minutes": 40})
    assert project(c, pid)["presentation_status"] == "analyzed"  # not left as `failed`


def test_approving_again_and_rebuilding_the_same_prompt_keep_the_presentation(client):
    pid, _ = generated(client)
    assert client.post(f"/projects/{pid}/plan/approve").status_code == 200
    assert project(client, pid)["presentation_status"] == "completed"
    again = client.post(f"/projects/{pid}/prompt")  # identical inputs -> identical text
    assert again.status_code == 200
    pr = project(client, pid)
    assert pr["presentation_status"] == "completed" and pr["has_presentation"] is True
    assert client.get(f"/projects/{pid}/presentation/download").status_code == 200


def test_a_different_prompt_withdraws_the_presentation(client):
    pid, _ = generated(client)
    assert client.post(f"/projects/{pid}/enrich").status_code == 200  # rebuilds the content
    pr = project(client, pid)
    assert pr["has_presentation"] is False and pr["presentation_status"] == "enriched"
    assert client.post(f"/projects/{pid}/prompt").status_code == 200
    assert client.post(f"/projects/{pid}/presentation").json()["state"] == "completed"


# ------------------------------------------------------------------ small pieces
@pytest.mark.parametrize("title, ext, mock, expected", [
    ("MQTT 입문", ".pptx", True, "MQTT 입문_MOCK.pptx"),
    ("MQTT 입문", ".pdf", False, "MQTT 입문.pdf"),
    ('a/b\\c:d*e?"f<g>h|', ".pptx", False, "a_b_c_d_e_f_g_h.pptx"),
    ("", ".pptx", False, "presentation.pptx"),
    (None, ".pptx", True, "presentation_MOCK.pptx"),
    ("..", ".pptx", False, "presentation.pptx"),
])
def test_download_names_are_safe(title, ext, mock, expected):
    assert _download_name(title, ext, mock) == expected

