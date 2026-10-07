"""STAGE 7 providers: the interface, MockPresentationProvider, provider selection, and the rule that
Genspark-specific code stays in providers/genspark.py. (The Genspark CLI provider itself: STAGE 8 tests.)"""

import io
import re
from pathlib import Path

import pytest
from pptx import Presentation

import app
from app.config import GENSPARK_MODES, Settings, load_settings, parse_genspark_mode
from app.errors import PresentationGenerationFailed, PresentationProviderUnavailable
from app.models.presentation import JobState, PresentationRequest, SlideBrief
from app.providers.base import PresentationProvider
from app.providers.factory import create_provider
from app.providers.genspark import GensparkProvider
from app.providers.mock import MockPresentationProvider

APP_DIR = Path(app.__file__).parent


def request(n=4):
    slides = [SlideBrief(slide_number=i, title=f"슬라이드 제목 {i}", slide_type="concept") for i in range(1, n + 1)]
    return PresentationRequest(
        project_id="a" * 32, title="MQTT 입문", final_prompt="prompt text", prompt_hash="f" * 64,
        slide_count=n, slides=slides, duration_minutes=60,
    )


# ------------------------------------------------------------------ the interface
def test_the_interface_has_the_four_operations_of_the_spec():
    for name in ("create_presentation", "get_status", "get_result", "download"):
        assert name in PresentationProvider.__abstractmethods__
    with pytest.raises(TypeError):
        PresentationProvider()  # abstract
    assert issubclass(MockPresentationProvider, PresentationProvider)
    assert issubclass(GensparkProvider, PresentationProvider)


# ------------------------------------------------------------------ MockPresentationProvider
def test_the_mock_renders_one_slide_per_planned_slide_in_order(tmp_path):
    p = MockPresentationProvider(tmp_path)
    st = p.create_presentation(request(5))
    assert st.state == JobState.completed and st.progress == 100 and st.job_id.startswith("mock-")
    res = p.get_result(st.job_id)
    f = p.download(st.job_id)
    assert res.slide_count == 5 and res.size_bytes == len(f.data) > 0
    assert f.file_name.endswith(".pptx") and "presentationml" in f.content_type
    prs = Presentation(io.BytesIO(f.data))  # a real, readable deck
    assert [s.shapes.title.text for s in prs.slides] == [f"슬라이드 제목 {i}" for i in range(1, 6)]
    for s in prs.slides:
        assert "MOCK" in " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
    assert "MOCK" in prs.core_properties.comments


def test_the_mock_is_marked_as_a_mock(tmp_path):
    p = MockPresentationProvider(tmp_path)
    assert p.is_mock is True and p.name == "mock"
    assert p.availability().configured is True


def test_a_slow_mock_job_moves_through_the_states(tmp_path):
    p = MockPresentationProvider(tmp_path, polls_until_done=2)
    st = p.create_presentation(request())
    assert st.state == JobState.queued and st.progress == 0
    with pytest.raises(PresentationGenerationFailed):
        p.get_result(st.job_id)  # not finished yet
    assert p.get_status(st.job_id).state == JobState.running
    assert p.get_status(st.job_id).state == JobState.completed
    assert p.get_result(st.job_id).slide_count == 4
    assert p.get_status(st.job_id).state == JobState.completed  # stays completed


def test_progress_grows(tmp_path):
    p = MockPresentationProvider(tmp_path, polls_until_done=4)
    st = p.create_presentation(request())
    seen = [st.progress] + [p.get_status(st.job_id).progress for _ in range(4)]
    assert seen == sorted(seen) and seen[0] == 0 and seen[-1] == 100


def test_the_mock_can_fail_while_rendering(tmp_path):
    p = MockPresentationProvider(tmp_path, fail_at="render")
    st = p.create_presentation(request())
    assert st.state == JobState.failed and st.message
    with pytest.raises(PresentationGenerationFailed):
        p.download(st.job_id)


def test_the_mock_can_refuse_to_start(tmp_path):
    with pytest.raises(PresentationProviderUnavailable):
        MockPresentationProvider(tmp_path, fail_at="create").create_presentation(request())
    with pytest.raises(ValueError):
        MockPresentationProvider(tmp_path, fail_at="whenever")


def test_mock_jobs_survive_a_restart(tmp_path):
    st = MockPresentationProvider(tmp_path, polls_until_done=1).create_presentation(request())
    again = MockPresentationProvider(tmp_path)  # a new process
    assert again.get_status(st.job_id).state == JobState.completed
    assert again.download(st.job_id).data


@pytest.mark.parametrize("job_id", ["../x", "mock-../../etc", "mock-123", "", "mock-" + "g" * 32, "x" * 40])
def test_bad_job_ids_are_rejected(tmp_path, job_id):
    p = MockPresentationProvider(tmp_path)
    for call in (p.get_status, p.get_result, p.download):
        with pytest.raises(PresentationGenerationFailed):
            call(job_id)


def test_an_unknown_but_wellformed_job_id_is_a_failure_not_a_crash(tmp_path):
    with pytest.raises(PresentationGenerationFailed):
        MockPresentationProvider(tmp_path).get_status("mock-" + "0" * 32)


# ------------------------------------------------------------------ selection and configuration
def test_the_default_provider_is_the_mock(tmp_path):
    p = create_provider(Settings(data_dir=tmp_path, max_upload_mb=1, cors_origins=()))
    assert isinstance(p, MockPresentationProvider)


def test_genspark_mode_selects_the_genspark_provider(tmp_path):
    p = create_provider(Settings(data_dir=tmp_path, max_upload_mb=1, cors_origins=(), genspark_mode="genspark"))
    assert isinstance(p, GensparkProvider)


@pytest.mark.parametrize("raw, expected", [(None, "mock"), ("", "mock"), ("MOCK", "mock"), (" genspark ", "genspark"), ("live", "mock"), ("real", "mock")])
def test_genspark_mode_parsing(raw, expected):
    assert parse_genspark_mode(raw) == expected
    assert expected in GENSPARK_MODES


def test_load_settings_reads_genspark_mode(monkeypatch):
    monkeypatch.setenv("GENSPARK_MODE", "genspark")
    assert load_settings().genspark_mode == "genspark"
    monkeypatch.delenv("GENSPARK_MODE")
    assert load_settings().genspark_mode == "mock"


# ------------------------------------------------------------------ architecture
def _py_files():
    return [p for p in APP_DIR.rglob("*.py") if "__pycache__" not in p.parts]


def test_genspark_specific_code_lives_only_in_the_provider_module():
    """Nothing but providers/genspark.py and the factory that chooses it may import GensparkProvider."""
    importers = {
        p.relative_to(APP_DIR).as_posix()
        for p in _py_files()
        if re.search(r"^\s*(?:from\s+\S*genspark\S*\s+import|import\s+\S*genspark)", p.read_text(encoding="utf-8"), re.M)
    }
    assert importers == {"providers/factory.py"}, importers


def test_the_lecture_engine_does_not_depend_on_any_concrete_provider():
    engine = [
        p for p in _py_files()
        if p.relative_to(APP_DIR).parts[0] in ("models", "llm", "storage")
        or p.relative_to(APP_DIR).as_posix().startswith("services/")
    ]
    assert engine
    for p in engine:
        text = p.read_text(encoding="utf-8")
        for m in re.finditer(r"^\s*(?:from|import)\s+([.\w]*providers[.\w]*)", text, re.M):
            module = m.group(1)
            ok = p.name == "presentation_service.py" and module.endswith("providers.base")
            assert ok, f"{p.name} imports {module}"


def test_the_provider_modules_do_not_reach_into_the_engine():
    """A provider needs only the models, the errors and its base."""
    for p in (APP_DIR / "providers").glob("*.py"):
        for m in re.finditer(r"^\s*from\s+(\.\.[\w.]*)\s+import", p.read_text(encoding="utf-8"), re.M):
            assert m.group(1) in ("..models.presentation", "..errors"), f"{p.name} imports {m.group(1)}"
